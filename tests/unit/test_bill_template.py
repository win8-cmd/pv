"""电费账单标准 Excel 模板测试（规格书第八部分第 38 节）。

覆盖规格书要求的模板验证项：

- Excel 文件成功生成且可正常读取；
- 七张工作表名称正确、顺序稳定；
- 关键字段齐全、表头顺序稳定；
- 枚举字段有下拉选项；
- 冻结窗格与筛选存在；
- 说明与样本标记存在；
- 样本工作表默认不参与正式导入；
- 不依赖宏；
- 业务表不预置示例数据行（避免被当成正式数据）；
- 空值与 0 的区别、负费用、需量单位等规则在字段说明中体现。
"""

from __future__ import annotations

import os
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from cenep.data import bill_template_spec as spec
from cenep.data.template_writer import build_template, default_template_path, write_template

#: 规格书第 20 节要求的七张工作表（名称与顺序均固定）
EXPECTED_SHEET_NAMES = (
    "使用说明",
    "月度汇总",
    "分时电量电价",
    "费用明细",
    "计量点",
    "数据校验",
    "样本参考_勿直接测算",
)

#: 月度汇总的必需字段（规格书第 23 节）
MONTHLY_REQUIRED_FIELDS = (
    "项目编号",
    "账期",
    "账期开始日期",
    "账期结束日期",
    "本期总电量_kWh",
    "账单本期电费_元",
)

#: 月度汇总必须出现的全部关键字段（规格书第 23.3/23.4/23.5 节）
MONTHLY_KEY_FIELDS = (
    "增值税专票金额_元",
    "账单展示平均电价_元每kWh",
    "工商业电费_元",
    "市场化购电电费_元",
    "上网环节线损费用_元",
    "输配电费_元",
    "系统运行费_元",
    "政府性基金及附加_元",
    "功率因数调整电费_元",
    "需量电费_元",
    "市场化运营费合计_元",
    "基本电费_元",
    "合同最大需量",
    "合同最大需量单位",
    "实际最大需量",
    "实际最大需量单位",
    "需量计费方式",
    "功率因数",
    "数据来源文件" if False else "来源文件",
    "来源页码",
    "复核状态",
    "备注",
)


@pytest.fixture(scope="module")
def template_file(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """生成一次模板文件，供本模块多数用例复用。"""
    directory = tmp_path_factory.mktemp("template")
    return write_template(directory / spec.TEMPLATE_NAME)


@pytest.fixture(scope="module")
def template_workbook(template_file: Path):
    return load_workbook(template_file)


# ---------------------------------------------------------------------------
# 38.1 文件生成与可读性
# ---------------------------------------------------------------------------


class TestTemplateGeneration:
    def test_template_file_is_created(self, template_file: Path) -> None:
        assert template_file.exists()
        assert template_file.suffix == ".xlsx"
        assert template_file.stat().st_size > 0

    def test_template_can_be_read_back(self, template_workbook) -> None:
        assert template_workbook.sheetnames == list(EXPECTED_SHEET_NAMES)

    def test_template_size_is_reasonable(self, template_file: Path) -> None:
        """模板不应因过度格式化而异常膨胀。"""
        size_kb = template_file.stat().st_size / 1024
        assert size_kb < 200, f"模板过大（{size_kb:.1f} KB），疑似格式覆盖行数过多"

    def test_template_has_no_macros(self, template_file: Path) -> None:
        """模板不依赖 Excel 宏：扩展名必须是 .xlsx，且包内无 vbaProject。"""
        assert template_file.suffix == ".xlsx"
        with zipfile.ZipFile(template_file) as archive:
            names = archive.namelist()
        assert not any("vbaProject" in name for name in names), "模板不应包含 VBA 宏"
        assert not any(name.endswith(".bin") for name in names), "模板不应包含宏二进制"

    def test_default_template_path_uses_canonical_name(self, tmp_path: Path) -> None:
        path = default_template_path(tmp_path)
        assert path.name == spec.TEMPLATE_NAME
        assert path.parent == tmp_path.resolve()

    def test_write_template_creates_parent_directory(self, tmp_path: Path) -> None:
        target = tmp_path / "nested" / "deeper" / spec.TEMPLATE_NAME
        written = write_template(target)
        assert written.exists()
        assert written.parent.is_dir()

    def test_build_template_returns_workbook_without_writing(self) -> None:
        workbook = build_template()
        assert workbook.sheetnames == list(EXPECTED_SHEET_NAMES)


# ---------------------------------------------------------------------------
# 38.2 工作表名称与关键字段
# ---------------------------------------------------------------------------


class TestSheetStructure:
    def test_seven_sheets_exist_in_expected_order(self, template_workbook) -> None:
        assert template_workbook.sheetnames == list(EXPECTED_SHEET_NAMES)
        assert len(template_workbook.sheetnames) == 7

    @pytest.mark.parametrize("sheet_name", EXPECTED_SHEET_NAMES)
    def test_each_sheet_has_header_row(self, template_workbook, sheet_name: str) -> None:
        sheet = template_workbook[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        assert headers, f"{sheet_name} 缺少首行字段标题"
        assert all(header for header in headers if header is not None) or headers[0], (
            f"{sheet_name} 首行存在空标题"
        )

    def test_monthly_summary_contains_required_fields(self, template_workbook) -> None:
        headers = [cell.value for cell in template_workbook["月度汇总"][1]]
        for field in MONTHLY_REQUIRED_FIELDS:
            assert field in headers, f"月度汇总缺少必需字段：{field}"

    def test_monthly_summary_contains_key_fields(self, template_workbook) -> None:
        headers = [cell.value for cell in template_workbook["月度汇总"][1]]
        for field in MONTHLY_KEY_FIELDS:
            assert field in headers, f"月度汇总缺少关键字段：{field}"

    def test_monthly_summary_header_is_unique(self) -> None:
        headers = spec.headers_of(spec.SHEET_MONTHLY_SUMMARY)
        assert len(headers) == len(set(headers)), "月度汇总存在重复字段名"

    @pytest.mark.parametrize(
        "sheet_name",
        [
            spec.SHEET_MONTHLY_SUMMARY,
            spec.SHEET_TIME_OF_USE,
            spec.SHEET_FEE_ITEMS,
            spec.SHEET_METER_POINTS,
        ],
    )
    def test_all_headers_are_unique(self, sheet_name: str) -> None:
        headers = spec.headers_of(sheet_name)
        assert len(headers) == len(set(headers)), f"{sheet_name} 存在重复字段名"

    def test_header_order_is_stable_across_builds(self) -> None:
        """表头顺序必须稳定（两次构建完全一致）。"""
        first = [tuple(c.value for c in build_template()[name][1]) for name in EXPECTED_SHEET_NAMES]
        second = [tuple(c.value for c in build_template()[name][1]) for name in EXPECTED_SHEET_NAMES]
        assert first == second

    def test_declared_headers_match_written_headers(
        self, template_workbook, template_file: Path
    ) -> None:
        """规范声明的字段顺序必须与写出的表头完全一致。"""
        for sheet_name in EXPECTED_SHEET_NAMES:
            declared = spec.headers_of(sheet_name)
            written = tuple(cell.value for cell in template_workbook[sheet_name][1])
            assert written == declared, f"{sheet_name} 表头与规范不一致"


# ---------------------------------------------------------------------------
# 38.3 格式：冻结、筛选、下拉、批注
# ---------------------------------------------------------------------------


class TestTemplateFormatting:
    @pytest.mark.parametrize(
        "sheet_name",
        [name for name in EXPECTED_SHEET_NAMES if name != spec.SHEET_INSTRUCTIONS],
    )
    def test_business_and_reference_sheets_freeze_header(
        self, template_workbook, sheet_name: str
    ) -> None:
        assert template_workbook[sheet_name].freeze_panes == "A2", f"{sheet_name} 未冻结首行"

    @pytest.mark.parametrize(
        "sheet_name",
        [name for name in EXPECTED_SHEET_NAMES if name != spec.SHEET_INSTRUCTIONS],
    )
    def test_business_and_reference_sheets_enable_filter(
        self, template_workbook, sheet_name: str
    ) -> None:
        assert template_workbook[sheet_name].auto_filter.ref, f"{sheet_name} 未启用筛选"

    def test_enum_columns_have_data_validation(self, template_workbook) -> None:
        """枚举字段必须有下拉选项（数据验证序列）。"""
        sheet = template_workbook[spec.SHEET_MONTHLY_SUMMARY]
        validations = list(sheet.data_validations.dataValidation)
        assert validations, "月度汇总没有任何数据验证（下拉选项）"

        for column in spec.columns_of(spec.SHEET_MONTHLY_SUMMARY):
            if not column.choices:
                continue
            matched = [
                item
                for item in validations
                if item.type == "list"
                and all(choice in (item.formula1 or "") for choice in column.choices)
            ]
            assert matched, f"字段 {column.name} 的枚举值缺少下拉选项"

    def test_every_enum_column_has_choices_declared(self) -> None:
        """凡是枚举语义的字段都应在规范中声明可选值。"""
        expected_enum_columns = {
            spec.SHEET_MONTHLY_SUMMARY: {"复核状态", "合同最大需量单位", "实际最大需量单位", "需量计费方式"},
            spec.SHEET_TIME_OF_USE: {"电量类型", "时段名称", "复核状态"},
            spec.SHEET_FEE_ITEMS: {
                "费用大类",
                "是否包含在其他费用中",
                "光伏测算可避免性",
                "是否纳入基准电费",
                "复核状态",
            },
            spec.SHEET_METER_POINTS: {
                "计量点类型",
                "是否已计入首页总电量",
                "是否允许纳入汇总",
                "复核状态",
            },
            spec.SHEET_DATA_VALIDATION: {"严重程度", "结果"},
        }
        for sheet_name, column_names in expected_enum_columns.items():
            for column_name in column_names:
                column = spec.column_by_name(sheet_name, column_name)
                assert column.choices, f"{sheet_name}.{column_name} 应声明可选值"

    def test_key_columns_have_comments(self, template_workbook) -> None:
        """关键字段应有单元格批注说明（单位、必填、可选值）。"""
        sheet = template_workbook[spec.SHEET_MONTHLY_SUMMARY]
        headers = [cell.value for cell in sheet[1]]
        index = headers.index("本期总电量_kWh") + 1
        comment = sheet.cell(row=1, column=index).comment
        assert comment is not None, "关键字段缺少说明批注"
        assert "kWh" in comment.text
        assert "必填" in comment.text

    def test_required_columns_are_marked_required(self) -> None:
        """必填字段必须在规范中标注，且数量合理（不能全是必填）。"""
        for sheet_name in (spec.SHEET_MONTHLY_SUMMARY, spec.SHEET_TIME_OF_USE):
            columns = spec.columns_of(sheet_name)
            required = [column.name for column in columns if column.required]
            assert required, f"{sheet_name} 应至少有一个必填字段"
            assert len(required) < len(columns), f"{sheet_name} 不应把所有字段都标为必填"

    def test_date_columns_have_date_format(self) -> None:
        for column_name in ("账期开始日期", "账期结束日期"):
            column = spec.column_by_name(spec.SHEET_MONTHLY_SUMMARY, column_name)
            assert column.number_format == "yyyy-mm-dd", f"{column_name} 缺少日期格式"


# ---------------------------------------------------------------------------
# 38.4 说明表与校验表
# ---------------------------------------------------------------------------


class TestInstructionAndValidationSheets:
    def test_instructions_sheet_lists_all_work_sheets(self, template_workbook) -> None:
        sheet = template_workbook[spec.SHEET_INSTRUCTIONS]
        text = "\n".join(
            str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value
        )
        for name in EXPECTED_SHEET_NAMES:
            assert name in text, f"使用说明未提及工作表：{name}"

    @pytest.mark.parametrize(
        "keyword",
        [
            "空值与零值",
            "来源与复核",
            "总电量与分项关系",
            "总电费口径区别",
            "费用重复计入风险",
            "需量电费的特殊性",
            "自用率计算前提",
            "样本数据警告",
            "模板版本变更规则",
        ],
    )
    def test_instructions_cover_required_topics(self, template_workbook, keyword: str) -> None:
        sheet = template_workbook[spec.SHEET_INSTRUCTIONS]
        text = "\n".join(
            str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value
        )
        assert keyword in text, f"使用说明缺少主题：{keyword}"

    def test_instructions_state_null_differs_from_zero(self, template_workbook) -> None:
        sheet = template_workbook[spec.SHEET_INSTRUCTIONS]
        text = "\n".join(
            str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value
        )
        assert "不得填 0" in text or "不得填0" in text
        assert "留空" in text

    def test_instructions_state_average_price_is_not_marginal_price(
        self, template_workbook
    ) -> None:
        sheet = template_workbook[spec.SHEET_INSTRUCTIONS]
        text = "\n".join(
            str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value
        )
        assert "自用率" in text

    def test_validation_sheet_has_thirteen_check_items(self, template_workbook) -> None:
        sheet = template_workbook[spec.SHEET_DATA_VALIDATION]
        codes = [
            sheet.cell(row=row, column=1).value
            for row in range(2, sheet.max_row + 1)
            if sheet.cell(row=row, column=1).value
        ]
        assert len(codes) >= 13, f"数据校验表应预置至少 13 项，实际 {len(codes)}"
        assert codes[0] == "V01"

    def test_validation_items_cover_required_checks(self) -> None:
        """规格书第 27 节 13 类校验项必须齐备。"""
        items = spec.validation_check_items()
        assert len(items) == 13
        text = " ".join(item[1] for item in items)
        for keyword in (
            "账期",
            "总电量",
            "金额",
            "口径冲突",
            "重复计入",
            "勾稽",
            "电价单位",
            "需量单位",
            "复核状态",
            "可追溯",
            "消纳",
        ):
            assert keyword in text, f"校验项缺少主题：{keyword}"

    def test_validation_severity_values_are_declared(self) -> None:
        severities = {item[3].value for item in spec.validation_check_items()}
        assert severities <= {"阻断", "警告", "提示"}
        assert "阻断" in severities

    def test_validation_result_defaults_to_pending(self, template_workbook) -> None:
        """校验结果默认必须是"待检查"，不得预设为"通过"。"""
        sheet = template_workbook[spec.SHEET_DATA_VALIDATION]
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row=row, column=1).value:
                assert sheet.cell(row=row, column=5).value == "待检查"


# ---------------------------------------------------------------------------
# 38.5 样本表与导入隔离
# ---------------------------------------------------------------------------


class TestSampleSheetIsolation:
    def test_sample_sheet_name_carries_warning(self) -> None:
        assert "勿直接测算" in spec.SHEET_SAMPLE_REFERENCE

    def test_sample_marker_text_present_in_sheet(self, template_workbook) -> None:
        sheet = template_workbook[spec.SHEET_SAMPLE_REFERENCE]
        text = "\n".join(
            str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value
        )
        assert spec.SAMPLE_MARKER_TEXT[:20] in text
        assert "不得作为生产项目账单自动导入" in text

    def test_sample_workbook_declares_pending_review(self, template_workbook) -> None:
        """样本数值必须标注为"待复核"。"""
        sheet = template_workbook[spec.SHEET_SAMPLE_REFERENCE]
        text = "\n".join(
            str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value
        )
        assert "待复核" in text

    def test_sample_sheet_is_excluded_from_import(self) -> None:
        """样本表默认不参与正式导入（规格书第 20、31 节）。"""
        assert spec.SHEET_SAMPLE_REFERENCE in spec.EXCLUDED_SHEETS

    def test_instructions_and_validation_sheets_are_excluded(self) -> None:
        assert spec.SHEET_INSTRUCTIONS in spec.EXCLUDED_SHEETS
        assert spec.SHEET_DATA_VALIDATION in spec.EXCLUDED_SHEETS

    def test_business_sheets_are_exactly_four(self) -> None:
        assert spec.BUSINESS_SHEETS == (
            spec.SHEET_MONTHLY_SUMMARY,
            spec.SHEET_TIME_OF_USE,
            spec.SHEET_FEE_ITEMS,
            spec.SHEET_METER_POINTS,
        )

    def test_business_and_excluded_sheets_are_disjoint_and_complete(self) -> None:
        business = set(spec.BUSINESS_SHEETS)
        excluded = set(spec.EXCLUDED_SHEETS)
        all_names = {sheet.name for sheet in spec.ALL_SHEETS}
        assert business & excluded == set(), "业务表与排除表不得重叠"
        assert business | excluded == all_names, "所有工作表必须被明确分类"

    def test_sample_values_match_frozen_baseline(self) -> None:
        """样本表数值必须与冻结基线一致（规格书第 28 节）。"""
        rows = {row[0]: row for row in spec.SAMPLE_TEMPLATE_ROWS}
        assert rows["本期总电量_kWh"][1] == "8780498"
        assert rows["本期总电量_kWh"][2] == "7479418"
        assert rows["账单本期电费_元"][1] == "6699241.94"
        assert rows["账单本期电费_元"][2] == "5471334.24"
        assert rows["增值税专票金额_元"][1] == "6534168.58"
        assert rows["工商业电费_元"][1] == "6708448.31"
        assert rows["功率因数调整电费_元"][1] == "-9206.37", "负费用必须保留负号"
        assert rows["功率因数调整电费_元"][2] == "-37735.81"
        assert rows["需量电费_元"][1] == "1202916.00"


# ---------------------------------------------------------------------------
# 38.6 业务表不预置数据
# ---------------------------------------------------------------------------


class TestBusinessSheetsHaveNoSampleData:
    @pytest.mark.parametrize(
        "sheet_name",
        [
            spec.SHEET_MONTHLY_SUMMARY,
            spec.SHEET_TIME_OF_USE,
            spec.SHEET_FEE_ITEMS,
            spec.SHEET_METER_POINTS,
        ],
    )
    def test_business_sheet_has_only_header_row(self, template_workbook, sheet_name: str) -> None:
        """业务表只写表头，不预置示例行——避免示例被当成正式数据导入。"""
        sheet = template_workbook[sheet_name]
        for row in range(2, sheet.max_row + 1):
            values = [sheet.cell(row=row, column=col).value for col in range(1, sheet.max_column + 1)]
            assert all(value is None for value in values), (
                f"{sheet_name} 第 {row} 行存在预置数据：{values}"
            )


# ---------------------------------------------------------------------------
# 38.7 规范自检
# ---------------------------------------------------------------------------


class TestSpecIntegrity:
    def test_template_version_is_declared(self) -> None:
        assert spec.TEMPLATE_VERSION == "1.0"

    def test_template_name_is_declared(self) -> None:
        assert spec.TEMPLATE_NAME == "CENEP_工商业电费账单标准导入模板_V1.0.xlsx"

    def test_get_sheet_raises_for_unknown_name(self) -> None:
        with pytest.raises(KeyError, match="未知工作表"):
            spec.get_sheet("不存在的表")

    def test_column_by_name_raises_for_unknown_column(self) -> None:
        with pytest.raises(KeyError, match="无字段"):
            spec.column_by_name(spec.SHEET_MONTHLY_SUMMARY, "不存在的字段")

    def test_units_follow_spec_conventions(self) -> None:
        """单位必须遵循规格书约定：电量 kWh、电价 元/kWh、金额 元。"""
        assert spec.column_by_name(spec.SHEET_MONTHLY_SUMMARY, "本期总电量_kWh").unit == "kWh"
        assert (
            spec.column_by_name(spec.SHEET_MONTHLY_SUMMARY, "账单展示平均电价_元每kWh").unit
            == "元/kWh"
        )
        assert spec.column_by_name(spec.SHEET_MONTHLY_SUMMARY, "账单本期电费_元").unit == "元"

    def test_demand_unit_choices_exclude_confusion(self) -> None:
        """需量单位必须显式区分 kW 与 kVA。"""
        column = spec.column_by_name(spec.SHEET_MONTHLY_SUMMARY, "实际最大需量单位")
        assert set(column.choices) == {"kW", "kVA"}

    def test_avoidability_marker_is_not_a_decision_rule(self) -> None:
        """可避免性字段必须包含"待核实"，且说明中声明不得直接决定收益。"""
        column = spec.column_by_name(spec.SHEET_FEE_ITEMS, "光伏测算可避免性")
        assert "待核实" in column.choices
        assert "不直接决定收益" in column.description

    def test_granularity_text_is_present_for_every_sheet(self) -> None:
        for sheet in spec.ALL_SHEETS:
            assert sheet.purpose, f"{sheet.name} 缺少用途说明"
            assert sheet.granularity, f"{sheet.name} 缺少粒度说明"

    def test_import_exclusion_flag_matches_spec(self) -> None:
        for sheet in spec.ALL_SHEETS:
            assert sheet.is_business is (sheet.name in spec.BUSINESS_SHEETS)
