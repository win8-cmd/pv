"""账单应用服务测试（人工录入、确认闸门、汇总、复制上月）。

覆盖规格书第 33 节"人工录入与编辑"与第 30 节数据纪律：

- 人工录入：必填校验、空值**不补 0**、非法数值被拒绝；
- 编辑与删除：整体替换语义、已确认账单删除需显式授权；
- 复制上月：继承静态信息、**金额与电量必须显式给出**（不得顺手带过去）、
  复核与确认状态重置；
- 确认闸门：**未确认账单不得进入正式测算**（``usable_bills`` 是唯一出口）；
- 样本数据不得确认；
- 汇总与筛选；
- 人工录入与 Excel 导入共用同一套校验口径。
"""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from cenep.application.bill_service import (
    BillService,
    BillSummary,
    ManualBillInput,
)
from cenep.data.bill_import_service import ImportDecision
from cenep.data.bill_template_spec import (
    SHEET_MONTHLY_SUMMARY,
    SHEET_TIME_OF_USE,
    TEMPLATE_NAME,
    headers_of,
)
from cenep.data.template_writer import write_template
from cenep.domain.bill_models import ReviewStatus
from cenep.domain.exceptions import PersistenceError, ValidationError
from cenep.domain.models import Project
from cenep.infrastructure.database import Database, ProjectRepository


@pytest.fixture
def service(tmp_path: Path):
    db = Database(tmp_path / "cenep.db")
    db.initialize_schema()
    project = ProjectRepository(db).add(Project(name="账单服务测试项目"))
    svc = BillService(db)
    yield svc, project.project_id
    db.close()


def _sep(**overrides) -> ManualBillInput:
    values = dict(
        billing_period="2025-09",
        period_start="2025-09-01",
        period_end="2025-09-30",
        account_number="1234567890",
        customer_name="示例用户",
        total_energy_kwh="8780498",
        total_amount_yuan="6699241.94",
        power_factor_adjustment_fee_yuan="-9206.37",
        actual_demand="30844",
        actual_demand_unit="kW",
        source_file="2025-09账单.pdf",
        source_page="1",
        review_status="已复核",
    )
    values.update(overrides)
    return ManualBillInput(**values)


# ---------------------------------------------------------------------------
# 人工录入
# ---------------------------------------------------------------------------


class TestManualEntry:
    def test_create_bill_persists(self, service) -> None:
        svc, project_id = service
        bill = svc.create_bill(project_id, _sep())
        assert bill.billing_period == "2025-09"
        assert bill.project_id == project_id
        stored = svc.list_bills(project_id)
        assert len(stored) == 1
        assert stored[0].bill_id == bill.bill_id

    def test_amounts_are_decimal(self, service) -> None:
        svc, project_id = service
        bill = svc.create_bill(project_id, _sep())
        assert isinstance(bill.total_amount_yuan, Decimal)
        assert bill.total_amount_yuan == Decimal("6699241.94")
        assert isinstance(bill.total_energy_kwh, Decimal)

    def test_blank_field_stays_none_not_zero(self, service) -> None:
        """空白输入必须保持未提供，**绝不补 0**。"""
        svc, project_id = service
        bill = svc.create_bill(
            project_id,
            ManualBillInput(
                billing_period="2025-09",
                account_number="1",
                total_energy_kwh="100",
                total_amount_yuan="200",
            ),
        )
        assert bill.vat_invoice_amount_yuan is None
        assert bill.demand_charge_yuan is None
        assert bill.actual_demand is None
        assert bill.vat_invoice_amount_yuan != Decimal(0)

    def test_whitespace_only_is_treated_as_blank(self, service) -> None:
        svc, project_id = service
        bill = svc.create_bill(project_id, _sep(vat_invoice_amount_yuan="   "))
        assert bill.vat_invoice_amount_yuan is None

    def test_explicit_zero_is_kept(self, service) -> None:
        svc, project_id = service
        bill = svc.create_bill(project_id, _sep(demand_charge_yuan="0"))
        assert bill.demand_charge_yuan == Decimal(0)
        assert bill.demand_charge_yuan is not None

    def test_negative_fee_accepted(self, service) -> None:
        svc, project_id = service
        bill = svc.create_bill(project_id, _sep())
        assert bill.power_factor_adjustment_fee_yuan == Decimal("-9206.37")

    def test_invalid_number_is_rejected(self, service) -> None:
        svc, project_id = service
        with pytest.raises(ValidationError, match="不是有效数值"):
            svc.create_bill(project_id, _sep(total_energy_kwh="约八百八十万"))

    def test_invalid_date_is_rejected(self, service) -> None:
        svc, project_id = service
        with pytest.raises(ValidationError, match="不是有效日期"):
            svc.create_bill(project_id, _sep(period_start="2025年9月1日"))

    def test_negative_total_energy_is_rejected(self, service) -> None:
        svc, project_id = service
        with pytest.raises(ValidationError, match="不得为负数"):
            svc.create_bill(project_id, _sep(total_energy_kwh="-1"))

    def test_invalid_demand_unit_is_rejected(self, service) -> None:
        svc, project_id = service
        with pytest.raises(ValidationError, match="kW 或 kVA"):
            svc.create_bill(project_id, _sep(actual_demand_unit="MW"))

    def test_invalid_review_status_is_rejected(self, service) -> None:
        svc, project_id = service
        with pytest.raises(ValidationError, match="复核状态非法"):
            svc.create_bill(project_id, _sep(review_status="随便写的"))

    def test_dict_payload_is_accepted(self, service) -> None:
        svc, project_id = service
        bill = svc.create_bill(
            project_id,
            {"billing_period": "2025-09", "total_energy_kwh": "1", "total_amount_yuan": "2"},
        )
        assert bill.total_energy_kwh == Decimal("1")

    def test_duplicate_business_key_is_rejected(self, service) -> None:
        """同一 项目+户号+账期 不允许两条。"""
        svc, project_id = service
        svc.create_bill(project_id, _sep())
        with pytest.raises(PersistenceError):
            svc.create_bill(project_id, _sep())

    def test_same_account_different_period_is_allowed(self, service) -> None:
        svc, project_id = service
        svc.create_bill(project_id, _sep())
        second = svc.create_bill(project_id, _sep(billing_period="2025-10"))
        assert second.billing_period == "2025-10"
        assert len(svc.list_bills(project_id)) == 2


# ---------------------------------------------------------------------------
# 编辑与删除
# ---------------------------------------------------------------------------


class TestEditAndDelete:
    def test_update_replaces_fields(self, service) -> None:
        svc, project_id = service
        bill = svc.create_bill(project_id, _sep())
        updated = svc.update_bill(
            bill.bill_id,
            _sep(total_amount_yuan="7000000.00", note="已按正式账单更正"),
        )
        assert updated.total_amount_yuan == Decimal("7000000.00")
        assert updated.note == "已按正式账单更正"
        assert len(svc.list_bills(project_id)) == 1, "更新不应新增记录"

    def test_update_resets_unspecified_fields(self, service) -> None:
        """整体替换语义：未提供的字段会被清空（避免残留旧值）。"""
        svc, project_id = service
        bill = svc.create_bill(project_id, _sep())
        updated = svc.update_bill(
            bill.bill_id,
            ManualBillInput(
                billing_period="2025-09",
                total_energy_kwh="100",
                total_amount_yuan="200",
            ),
        )
        assert updated.power_factor_adjustment_fee_yuan is None
        assert updated.actual_demand is None

    def test_update_missing_bill_raises(self, service) -> None:
        svc, _ = service
        with pytest.raises(ValidationError, match="不存在"):
            svc.update_bill("no-such-bill", _sep())

    def test_update_resets_confirmation(self, service) -> None:
        """修改后确认状态必须重置：改过的数据需要重新确认。"""
        svc, project_id = service
        bill = svc.create_bill(project_id, _sep())
        svc.confirm_bill(bill.bill_id)
        assert svc.get_bill(bill.bill_id).is_confirmed is True

        updated = svc.update_bill(
            bill.bill_id, _sep(total_amount_yuan="1")
        )
        assert updated.is_confirmed is False, "修改后必须重新确认"

    def test_delete_removes_bill(self, service) -> None:
        svc, project_id = service
        bill = svc.create_bill(project_id, _sep())
        svc.delete_bill(bill.bill_id)
        assert svc.list_bills(project_id) == []

    def test_delete_confirmed_requires_authorisation(self, service) -> None:
        svc, project_id = service
        bill = svc.create_bill(project_id, _sep())
        svc.confirm_bill(bill.bill_id)
        with pytest.raises(ValidationError, match="已确认"):
            svc.delete_bill(bill.bill_id)
        assert len(svc.list_bills(project_id)) == 1

    def test_delete_confirmed_with_authorisation(self, service) -> None:
        svc, project_id = service
        bill = svc.create_bill(project_id, _sep())
        svc.confirm_bill(bill.bill_id)
        svc.delete_bill(bill.bill_id, allow_confirmed=True)
        assert svc.list_bills(project_id) == []

    def test_delete_missing_bill_raises(self, service) -> None:
        svc, _ = service
        with pytest.raises(ValidationError, match="不存在"):
            svc.delete_bill("no-such-bill")


# ---------------------------------------------------------------------------
# 复制上月
# ---------------------------------------------------------------------------


class TestCopyFromPreviousMonth:
    def test_copy_inherits_static_information(self, service) -> None:
        svc, project_id = service
        source = svc.create_bill(project_id, _sep())
        copy = svc.copy_bill(
            source.bill_id,
            billing_period="2025-10",
            period_start="2025-10-01",
            period_end="2025-10-31",
            total_energy_kwh="7479418",
            total_amount_yuan="5471334.24",
        )
        assert copy.account_number == source.account_number
        assert copy.customer_name == source.customer_name
        assert copy.actual_demand_unit == "kW"

    def test_copy_does_not_carry_amounts_forward(self, service) -> None:
        """**关键安全规则**：不给新金额时不得继承上月数值。"""
        svc, project_id = service
        source = svc.create_bill(project_id, _sep())
        copy = svc.copy_bill(source.bill_id, billing_period="2025-10")
        assert copy.total_energy_kwh is None, "电量不得继承上月"
        assert copy.total_amount_yuan is None, "金额不得继承上月"
        assert copy.power_factor_adjustment_fee_yuan is None

    def test_copy_resets_review_and_confirmation(self, service) -> None:
        svc, project_id = service
        source = svc.create_bill(project_id, _sep())
        copy = svc.copy_bill(
            source.bill_id,
            billing_period="2025-10",
            total_energy_kwh="1",
            total_amount_yuan="2",
        )
        assert copy.review_status is ReviewStatus.PENDING, "复核状态应重置为待复核"
        assert copy.is_confirmed is False

    def test_copy_missing_source_raises(self, service) -> None:
        svc, _ = service
        with pytest.raises(ValidationError, match="来源账单不存在"):
            svc.copy_bill("no-such-bill", billing_period="2025-10")

    def test_copy_to_existing_period_is_rejected(self, service) -> None:
        svc, project_id = service
        source = svc.create_bill(project_id, _sep())
        svc.create_bill(project_id, _sep(billing_period="2025-10"))
        with pytest.raises(PersistenceError):
            svc.copy_bill(source.bill_id, billing_period="2025-10")


# ---------------------------------------------------------------------------
# 确认闸门
# ---------------------------------------------------------------------------


class TestConfirmationGate:
    def test_new_bill_is_not_usable_for_calculation(self, service) -> None:
        svc, project_id = service
        bill = svc.create_bill(project_id, _sep())
        assert bill.is_confirmed is False
        assert bill.is_usable_for_calculation is False

    def test_usable_bills_is_empty_before_confirmation(self, service) -> None:
        """未确认账单不得进入正式测算 —— 这是唯一的出口闸门。"""
        svc, project_id = service
        svc.create_bill(project_id, _sep())
        assert svc.usable_bills(project_id) == []

    def test_usable_bills_after_confirmation(self, service) -> None:
        svc, project_id = service
        bill = svc.create_bill(project_id, _sep())
        svc.confirm_bill(bill.bill_id, confirmed_by="tester")
        usable = svc.usable_bills(project_id)
        assert len(usable) == 1
        assert usable[0].is_usable_for_calculation is True

    def test_sample_bill_cannot_be_confirmed(self, service) -> None:
        svc, project_id = service
        bill = svc.create_bill(project_id, _sep(), is_sample=True)
        with pytest.raises(ValidationError, match="样本"):
            svc.confirm_bill(bill.bill_id)

    def test_sample_bill_never_usable(self, service) -> None:
        svc, project_id = service
        svc.create_bill(project_id, _sep(), is_sample=True)
        assert svc.usable_bills(project_id) == []

    def test_confirm_missing_bill_raises(self, service) -> None:
        svc, _ = service
        with pytest.raises(ValidationError, match="不存在"):
            svc.confirm_bill("no-such-bill")

    def test_pending_review_bill_can_be_confirmed(self, service) -> None:
        """"待复核"只是"还没复核"，用户人工核对后确认即可用于测算。

        否则会出现"必须先复核才能确认、但复核又要先确认"的死循环。
        """
        svc, project_id = service
        bill = svc.create_bill(project_id, _sep(review_status="待复核"))
        svc.confirm_bill(bill.bill_id)
        assert len(svc.usable_bills(project_id)) == 1

    def test_questionable_bill_is_never_usable(self, service) -> None:
        """"存在疑问"说明已知数据有问题，即使被确认也不得用于测算。"""
        svc, project_id = service
        bill = svc.create_bill(project_id, _sep(review_status="存在疑问"))
        svc.confirm_bill(bill.bill_id)
        assert svc.usable_bills(project_id) == [], "存在疑问的账单不得测算"
        assert len(svc.bills_with_unresolved_questions(project_id)) == 1


# ---------------------------------------------------------------------------
# 复核状态与筛选
# ---------------------------------------------------------------------------


class TestReviewStatus:
    def test_set_review_status(self, service) -> None:
        svc, project_id = service
        bill = svc.create_bill(project_id, _sep())
        svc.set_review_status(bill.bill_id, ReviewStatus.QUESTIONABLE)
        assert svc.get_bill(bill.bill_id).review_status is ReviewStatus.QUESTIONABLE

    def test_set_review_status_accepts_enum_member(self, service) -> None:
        """必须接受枚举成员本身（ReviewStatus 是 str 子类，易踩坑）。"""
        svc, project_id = service
        bill = svc.create_bill(project_id, _sep())
        svc.set_review_status(bill.bill_id, ReviewStatus.REVIEWED)
        assert svc.get_bill(bill.bill_id).review_status is ReviewStatus.REVIEWED

    def test_set_review_status_accepts_text(self, service) -> None:
        svc, project_id = service
        bill = svc.create_bill(project_id, _sep())
        svc.set_review_status(bill.bill_id, "已复核")
        assert svc.get_bill(bill.bill_id).review_status is ReviewStatus.REVIEWED

    def test_set_invalid_review_status_raises(self, service) -> None:
        svc, project_id = service
        bill = svc.create_bill(project_id, _sep())
        with pytest.raises(ValidationError, match="复核状态非法"):
            svc.set_review_status(bill.bill_id, "乱填")

    def test_filter_by_review_status(self, service) -> None:
        svc, project_id = service
        svc.create_bill(project_id, _sep())
        svc.create_bill(project_id, _sep(billing_period="2025-10", review_status="待复核"))
        assert len(svc.list_bills_by_review_status(project_id, ReviewStatus.REVIEWED)) == 1
        assert len(svc.list_bills_by_review_status(project_id, "待复核")) == 1

    def test_filter_by_period(self, service) -> None:
        svc, project_id = service
        svc.create_bill(project_id, _sep())
        svc.create_bill(project_id, _sep(billing_period="2025-10"))
        assert len(svc.list_bills(project_id, billing_period="2025-10")) == 1


# ---------------------------------------------------------------------------
# 汇总
# ---------------------------------------------------------------------------


class TestSummary:
    def test_empty_project_summary(self, service) -> None:
        svc, project_id = service
        summary = svc.summarize(project_id)
        assert summary.bill_count == 0
        assert summary.total_amount_yuan is None
        assert summary.all_confirmed is False

    def test_summary_totals(self, service) -> None:
        svc, project_id = service
        svc.create_bill(project_id, _sep())
        svc.create_bill(
            project_id,
            _sep(billing_period="2025-10", total_energy_kwh="7479418",
                 total_amount_yuan="5471334.24"),
        )
        summary = svc.summarize(project_id)
        assert summary.bill_count == 2
        assert summary.periods == ("2025-09", "2025-10")
        assert summary.total_energy_kwh == Decimal("16259916")
        assert summary.total_amount_yuan == Decimal("12170576.18")

    def test_summary_skips_missing_values(self, service) -> None:
        """缺失值不得当成 0 参与合计。"""
        svc, project_id = service
        svc.create_bill(
            project_id,
            ManualBillInput(billing_period="2025-09", account_number="1",
                            total_amount_yuan="100"),
        )
        summary = svc.summarize(project_id)
        assert summary.total_energy_kwh is None, "全部缺失时应为 None 而不是 0"
        assert summary.total_amount_yuan == Decimal("100")

    def test_summary_counts_confirmed(self, service) -> None:
        svc, project_id = service
        bill = svc.create_bill(project_id, _sep())
        svc.create_bill(project_id, _sep(billing_period="2025-10"))
        assert svc.summarize(project_id).confirmed_count == 0
        svc.confirm_bill(bill.bill_id)
        assert svc.summarize(project_id).confirmed_count == 1

    def test_summary_to_dict_is_serialisable(self, service) -> None:
        svc, project_id = service
        svc.create_bill(project_id, _sep())
        payload = svc.summarize(project_id).to_dict()
        assert payload["bill_count"] == 1
        assert isinstance(payload["total_amount_yuan"], str)
        assert payload["all_confirmed"] is False

    def test_summary_is_a_bill_summary(self, service) -> None:
        svc, project_id = service
        assert isinstance(svc.summarize(project_id), BillSummary)


# ---------------------------------------------------------------------------
# 与 Excel 导入共用校验口径
# ---------------------------------------------------------------------------


class TestSharedValidationWithImport:
    def test_manual_bill_passes_same_13_checks(self, service) -> None:
        svc, project_id = service
        svc.create_bill(project_id, _sep())
        report = svc.validate_project_bills(project_id)
        assert report.checked_bills == 1
        assert len(report.issues_for_check("V01")) == 1
        assert report.issues_for_check("V01")[0].result.value == "通过"

    def test_manual_and_imported_bills_validated_together(
        self, service, tmp_path: Path
    ) -> None:
        svc, project_id = service
        svc.create_bill(project_id, _sep())

        template = write_template(tmp_path / TEMPLATE_NAME)
        workbook = load_workbook(template)
        headers = headers_of(SHEET_MONTHLY_SUMMARY)
        october = {
            "项目编号": "P", "户号": "1234567890", "账期": "2025-10",
            "账期开始日期": "2025-10-01", "账期结束日期": "2025-10-31",
            "本期总电量_kWh": 7479418, "账单本期电费_元": 5471334.24,
            "来源文件": "2025-10账单.pdf", "来源页码": "1", "复核状态": "已复核",
        }
        for index, header in enumerate(headers, start=1):
            if header in october:
                workbook[SHEET_MONTHLY_SUMMARY].cell(row=2, column=index, value=october[header])
        tou_headers = headers_of(SHEET_TIME_OF_USE)
        for index, header in enumerate(tou_headers, start=1):
            values = {
                "项目编号": "P", "账期": "2025-10", "计量口径/计量点编号": "MP1",
                "电量类型": "购电", "时段名称": "平", "电量_kWh": 7479418,
            }
            if header in values:
                workbook[SHEET_TIME_OF_USE].cell(row=2, column=index, value=values[header])
        workbook.save(template)

        preview = svc.preview_import(template, project_id=project_id)
        assert preview.counts.total == 1
        svc.apply_import(preview, decision=ImportDecision.VALID_ONLY)

        assert len(svc.list_bills(project_id)) == 2
        report = svc.validate_project_bills(project_id)
        assert report.checked_bills == 2
        # 人工录入的账单来自 .pdf，导入的也来自 .pdf，均可追溯
        assert report.issues_for_check("V12")[0].result.value == "通过"

    def test_preview_does_not_write(self, service, tmp_path: Path) -> None:
        svc, project_id = service
        template = write_template(tmp_path / TEMPLATE_NAME)
        svc.preview_import(template, project_id=project_id)
        assert svc.list_bills(project_id) == []
