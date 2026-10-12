"""Excel 导入读取测试（规格书第 39 节导入测试）。

覆盖规格书第 39 节要求的全部情形：

- 正常文件；
- 空白可选字段；
- 缺失必填字段；
- 错误日期；
- 负费用；
- 负总电量；
- 中文列名别名；
- 列顺序变化；
- 未知额外列；
- 缺少必需工作表；
- 损坏文件；
- 重复账期 / 重复文件导入（本模块验证可识别所需的文件摘要与行数据）；
- 有效与无效记录混合；
- 用户取消导入 / 仅导入有效行（属阶段 3 服务层，此处验证数据层提供判定依据）；
- 重复记录的跳过、替换、保留逻辑（同上）。

数据层负责"把表读成结构化行"，因此本模块重点验证读取的正确性与容错边界。
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from cenep.data.bill_template_spec import (
    BUSINESS_SHEETS,
    SHEET_FEE_ITEMS,
    SHEET_INSTRUCTIONS,
    SHEET_METER_POINTS,
    SHEET_MONTHLY_SUMMARY,
    SHEET_SAMPLE_REFERENCE,
    SHEET_TIME_OF_USE,
    TEMPLATE_NAME,
    columns_of,
    headers_of,
)
from cenep.data.template_reader import (
    TemplateReadError,
    build_column_mapping,
    compute_file_hash,
    map_header,
    normalize_header,
    read_template,
)
from cenep.data.template_writer import write_template

#: 2025-09 冻结基线（规格书第 28 节）
SEP_ROW = {
    "项目编号": "PRJ-001",
    "用户名称": "示例用户",
    "户号": "1234567890",
    "账期": "2025-09",
    "账期开始日期": "2025-09-01",
    "账期结束日期": "2025-09-30",
    "本期总电量_kWh": 8780498,
    "账单本期电费_元": 6699241.94,
    "增值税专票金额_元": 6534168.58,
    "账单展示平均电价_元每kWh": 0.76297,
    "工商业电费_元": 6708448.31,
    "功率因数调整电费_元": -9206.37,
    "需量电费_元": 1202916.00,
    "来源文件": "2025-09账单.pdf",
    "来源页码": "1",
    "复核状态": "已复核",
}


def _fill_row(sheet, headers: tuple[str, ...], values: dict[str, object], row: int = 2) -> None:
    """按表头把一条数据写入指定行（仅写提供的字段）。"""
    for index, header in enumerate(headers, start=1):
        if header in values:
            sheet.cell(row=row, column=index, value=values[header])


def _make_template(path: Path) -> Path:
    return write_template(path)


def _append_monthly_row(path: Path, values: dict[str, object], row: int = 2) -> Path:
    workbook = load_workbook(path)
    _fill_row(workbook[SHEET_MONTHLY_SUMMARY], headers_of(SHEET_MONTHLY_SUMMARY), values, row)
    workbook.save(path)
    return path


@pytest.fixture
def template_path(tmp_path: Path) -> Path:
    return _make_template(tmp_path / TEMPLATE_NAME)


# ---------------------------------------------------------------------------
# 39.1 正常文件
# ---------------------------------------------------------------------------


class TestNormalFile:
    def test_read_generated_template(self, template_path: Path) -> None:
        result = read_template(template_path)
        assert result.file_name == TEMPLATE_NAME
        assert result.template_version == "1.0"
        assert result.total_business_rows == 0
        assert set(result.rows) == set(BUSINESS_SHEETS)

    def test_read_row_with_all_fields(self, template_path: Path) -> None:
        _append_monthly_row(template_path, SEP_ROW)
        result = read_template(template_path)
        rows = result.business_rows(SHEET_MONTHLY_SUMMARY)
        assert len(rows) == 1
        row = rows[0]
        assert row.row_number == 2
        assert row.text("项目编号") == "PRJ-001"
        assert row.text("账期") == "2025-09"

    def test_numeric_values_read_as_decimal(self, template_path: Path) -> None:
        _append_monthly_row(template_path, SEP_ROW)
        row = read_template(template_path).business_rows(SHEET_MONTHLY_SUMMARY)[0]
        assert row.decimal("本期总电量_kWh") == Decimal("8780498")
        assert row.decimal("账单本期电费_元") == Decimal("6699241.94")

    def test_dates_read_as_date_objects(self, template_path: Path) -> None:
        _append_monthly_row(template_path, SEP_ROW)
        row = read_template(template_path).business_rows(SHEET_MONTHLY_SUMMARY)[0]
        assert row.date("账期开始日期") == date(2025, 9, 1)
        assert row.date("账期结束日期") == date(2025, 9, 30)

    def test_multiple_rows_keep_excel_row_numbers(self, template_path: Path) -> None:
        _append_monthly_row(template_path, SEP_ROW, row=2)
        _append_monthly_row(template_path, {**SEP_ROW, "账期": "2025-10"}, row=4)  # 第 3 行留空
        rows = read_template(template_path).business_rows(SHEET_MONTHLY_SUMMARY)
        assert [row.row_number for row in rows] == [2, 4], "空行应被跳过但行号必须真实"

    def test_all_four_business_sheets_are_read(self, template_path: Path) -> None:
        workbook = load_workbook(template_path)
        _fill_row(
            workbook[SHEET_TIME_OF_USE],
            headers_of(SHEET_TIME_OF_USE),
            {"项目编号": "PRJ-001", "账期": "2025-09", "计量口径/计量点编号": "MP1", "电量_kWh": 100},
        )
        _fill_row(
            workbook[SHEET_FEE_ITEMS],
            headers_of(SHEET_FEE_ITEMS),
            {"项目编号": "PRJ-001", "账期": "2025-09", "费用项目名称": "输配电费", "金额_元": 2000662.04},
        )
        _fill_row(
            workbook[SHEET_METER_POINTS],
            headers_of(SHEET_METER_POINTS),
            {"项目编号": "PRJ-001", "账期": "2025-09", "计量点编号": "M1", "电量_kWh": 8516838},
        )
        workbook.save(template_path)

        result = read_template(template_path)
        assert len(result.business_rows(SHEET_TIME_OF_USE)) == 1
        assert len(result.business_rows(SHEET_FEE_ITEMS)) == 1
        assert len(result.business_rows(SHEET_METER_POINTS)) == 1


# ---------------------------------------------------------------------------
# 39.2 空白可选字段与空值语义
# ---------------------------------------------------------------------------


class TestBlankOptionalFields:
    def test_blank_optional_field_returns_none(self, template_path: Path) -> None:
        minimal = {
            "项目编号": "PRJ-001",
            "账期": "2025-09",
            "账期开始日期": "2025-09-01",
            "账期结束日期": "2025-09-30",
            "本期总电量_kWh": 100,
            "账单本期电费_元": 200,
        }
        _append_monthly_row(template_path, minimal)
        row = read_template(template_path).business_rows(SHEET_MONTHLY_SUMMARY)[0]
        assert row.raw("增值税专票金额_元") is None
        assert row.raw("需量电费_元") is None
        assert row.text("备注") is None

    def test_blank_is_not_zero(self, template_path: Path) -> None:
        minimal = {"项目编号": "P", "账期": "2025-09", "本期总电量_kWh": None}
        workbook = load_workbook(template_path)
        _fill_row(workbook[SHEET_MONTHLY_SUMMARY], headers_of(SHEET_MONTHLY_SUMMARY), minimal)
        workbook.save(template_path)
        row = read_template(template_path).business_rows(SHEET_MONTHLY_SUMMARY)[0]
        assert row.decimal("本期总电量_kWh") is None
        assert row.decimal("本期总电量_kWh") != Decimal(0)

    def test_explicit_zero_is_preserved(self, template_path: Path) -> None:
        _append_monthly_row(template_path, {**SEP_ROW, "需量电费_元": 0})
        row = read_template(template_path).business_rows(SHEET_MONTHLY_SUMMARY)[0]
        assert row.decimal("需量电费_元") == Decimal(0)
        assert row.decimal("需量电费_元") is not None

    def test_whitespace_only_cell_is_none(self, template_path: Path) -> None:
        _append_monthly_row(template_path, {**SEP_ROW, "备注": "   "})
        row = read_template(template_path).business_rows(SHEET_MONTHLY_SUMMARY)[0]
        assert row.text("备注") is None


# ---------------------------------------------------------------------------
# 39.3 缺失必填字段
# ---------------------------------------------------------------------------


class TestMissingRequiredFields:
    def test_missing_project_code_reports_required_field(self, template_path: Path) -> None:
        headers = list(headers_of(SHEET_MONTHLY_SUMMARY))
        headers.remove("项目编号")
        workbook = load_workbook(template_path)
        sheet = workbook[SHEET_MONTHLY_SUMMARY]
        for column in range(1, sheet.max_column + 1):
            sheet.cell(row=1, column=column, value=None)
        for index, header in enumerate(headers, start=1):
            sheet.cell(row=1, column=index, value=header)
        workbook.save(template_path)

        with pytest.raises(TemplateReadError) as excinfo:
            read_template(template_path)
        assert "项目编号" in str(excinfo.value) or any(
            "项目编号" in detail for detail in excinfo.value.details
        )

    def test_missing_required_fields_listed_in_mapping(self) -> None:
        header_row = [h for h in headers_of(SHEET_MONTHLY_SUMMARY) if h != "账期"]
        mapping = build_column_mapping(SHEET_MONTHLY_SUMMARY, header_row)
        assert "账期" in mapping.missing_required

    def test_required_fields_are_declared_in_spec(self) -> None:
        required = [c.name for c in columns_of(SHEET_MONTHLY_SUMMARY) if c.required]
        assert "项目编号" in required
        assert "账期" in required
        assert "本期总电量_kWh" in required


# ---------------------------------------------------------------------------
# 39.4 错误日期与负值
# ---------------------------------------------------------------------------


class TestInvalidValues:
    def test_invalid_date_raises_on_access(self, template_path: Path) -> None:
        _append_monthly_row(template_path, {**SEP_ROW, "账期开始日期": "2025年9月1日"})
        row = read_template(template_path).business_rows(SHEET_MONTHLY_SUMMARY)[0]
        with pytest.raises(ValueError, match="无法解析为日期"):
            row.date("账期开始日期")

    def test_text_date_formats_are_accepted(self, template_path: Path) -> None:
        _append_monthly_row(template_path, {**SEP_ROW, "账期开始日期": "2025/09/01"})
        row = read_template(template_path).business_rows(SHEET_MONTHLY_SUMMARY)[0]
        assert row.date("账期开始日期") == date(2025, 9, 1)

    def test_excel_datetime_cell_is_accepted(self, template_path: Path) -> None:
        _append_monthly_row(template_path, {**SEP_ROW, "账期开始日期": datetime(2025, 9, 1, 8, 0)})
        row = read_template(template_path).business_rows(SHEET_MONTHLY_SUMMARY)[0]
        assert row.date("账期开始日期") == date(2025, 9, 1)

    def test_negative_fee_is_preserved(self, template_path: Path) -> None:
        _append_monthly_row(template_path, SEP_ROW)
        row = read_template(template_path).business_rows(SHEET_MONTHLY_SUMMARY)[0]
        assert row.decimal("功率因数调整电费_元") == Decimal("-9206.37")
        assert row.decimal("功率因数调整电费_元") < 0

    def test_negative_total_energy_is_read_verbatim(self, template_path: Path) -> None:
        """读取层如实读出负值（是否合法由校验层判定，读取层不得篡改）。"""
        _append_monthly_row(template_path, {**SEP_ROW, "本期总电量_kWh": -100})
        row = read_template(template_path).business_rows(SHEET_MONTHLY_SUMMARY)[0]
        assert row.decimal("本期总电量_kWh") == Decimal("-100")

    def test_unparsable_number_raises(self, template_path: Path) -> None:
        _append_monthly_row(template_path, {**SEP_ROW, "本期总电量_kWh": "约八百八十万"})
        row = read_template(template_path).business_rows(SHEET_MONTHLY_SUMMARY)[0]
        with pytest.raises(ValueError, match="无法解析为数值"):
            row.decimal("本期总电量_kWh")

    def test_thousand_separator_text_is_accepted(self, template_path: Path) -> None:
        _append_monthly_row(template_path, {**SEP_ROW, "本期总电量_kWh": "8,780,498"})
        row = read_template(template_path).business_rows(SHEET_MONTHLY_SUMMARY)[0]
        assert row.decimal("本期总电量_kWh") == Decimal("8780498")


# ---------------------------------------------------------------------------
# 39.5 列名别名与列顺序
# ---------------------------------------------------------------------------


class TestHeaderMapping:
    @pytest.mark.parametrize(
        "alias, expected",
        [
            ("本期电量", "本期总电量_kWh"),
            ("总用电量", "本期总电量_kWh"),
            ("总电量", "本期总电量_kWh"),
            ("力调电费", "功率因数调整电费_元"),
            ("专票金额", "增值税专票金额_元"),
            ("发票金额", "增值税专票金额_元"),
            ("平均电价", "账单展示平均电价_元每kWh"),
            ("月份", "账期"),
            ("用户号", "户号"),
            ("页码", "来源页码"),
        ],
    )
    def test_chinese_aliases_are_mapped(self, alias: str, expected: str) -> None:
        assert map_header(SHEET_MONTHLY_SUMMARY, alias) == expected

    def test_header_without_unit_suffix_is_mapped(self) -> None:
        assert map_header(SHEET_MONTHLY_SUMMARY, "本期总电量") == "本期总电量_kWh"
        assert map_header(SHEET_MONTHLY_SUMMARY, "账单本期电费") == "账单本期电费_元"

    def test_unknown_header_returns_none(self) -> None:
        assert map_header(SHEET_MONTHLY_SUMMARY, "完全不相关的列") is None

    def test_blank_header_returns_none(self) -> None:
        assert map_header(SHEET_MONTHLY_SUMMARY, None) is None
        assert map_header(SHEET_MONTHLY_SUMMARY, "   ") is None

    def test_header_order_change_is_tolerated(self, template_path: Path) -> None:
        """打乱列顺序后仍应正确映射到字段。"""
        headers = list(headers_of(SHEET_MONTHLY_SUMMARY))
        shuffled = headers[::-1]
        workbook = load_workbook(template_path)
        sheet = workbook[SHEET_MONTHLY_SUMMARY]
        for column in range(1, sheet.max_column + 1):
            sheet.cell(row=1, column=column, value=None)
        for index, header in enumerate(shuffled, start=1):
            sheet.cell(row=1, column=index, value=header)
        _fill_row(sheet, tuple(shuffled), SEP_ROW, row=2)
        workbook.save(template_path)

        row = read_template(template_path).business_rows(SHEET_MONTHLY_SUMMARY)[0]
        assert row.text("项目编号") == "PRJ-001"
        assert row.decimal("本期总电量_kWh") == Decimal("8780498")
        assert row.decimal("账单本期电费_元") == Decimal("6699241.94")

    def test_unknown_extra_column_is_reported_not_dropped(self) -> None:
        header_row = list(headers_of(SHEET_MONTHLY_SUMMARY)) + ["额外无关列"]
        mapping = build_column_mapping(SHEET_MONTHLY_SUMMARY, header_row)
        unknown_headers = [header for _index, header in mapping.unknown_columns]
        assert "额外无关列" in unknown_headers

    def test_duplicate_header_keeps_first_and_reports_second(self) -> None:
        header_row = ["项目编号", "项目编号"]
        mapping = build_column_mapping(SHEET_MONTHLY_SUMMARY, header_row)
        assert mapping.index_of("项目编号") == 0
        assert any(header == "项目编号" for _index, header in mapping.unknown_columns)

    def test_normalize_header_strips_units_and_separators(self) -> None:
        assert normalize_header("本期总电量_kWh") == "本期总电量"
        assert normalize_header("账单本期电费_元") == "账单本期电费"
        assert normalize_header(" 平均  电价 ") == "平均电价"
        assert normalize_header(None) == ""


# ---------------------------------------------------------------------------
# 39.6 工作表缺失、损坏文件
# ---------------------------------------------------------------------------


class TestFileLevelErrors:
    def test_missing_required_sheet_raises(self, tmp_path: Path) -> None:
        path = tmp_path / "no_sheets.xlsx"
        workbook = Workbook()
        workbook.active.title = "随便一张表"
        workbook.save(path)
        with pytest.raises(TemplateReadError, match="缺少必需工作表"):
            read_template(path)

    def test_file_not_found_raises(self, tmp_path: Path) -> None:
        with pytest.raises(TemplateReadError, match="文件不存在"):
            read_template(tmp_path / "missing.xlsx")

    def test_empty_file_raises(self, tmp_path: Path) -> None:
        path = tmp_path / "empty.xlsx"
        path.write_bytes(b"")
        with pytest.raises(TemplateReadError, match="文件为空"):
            read_template(path)

    def test_corrupt_file_raises(self, tmp_path: Path) -> None:
        path = tmp_path / "corrupt.xlsx"
        path.write_bytes(b"this is definitely not a workbook")
        with pytest.raises(TemplateReadError, match="无法读取工作簿"):
            read_template(path)

    def test_directory_path_raises(self, tmp_path: Path) -> None:
        with pytest.raises(TemplateReadError, match="不是普通文件"):
            read_template(tmp_path)


# ---------------------------------------------------------------------------
# 39.7 排除表与样本隔离
# ---------------------------------------------------------------------------


class TestExcludedSheets:
    def test_excluded_sheets_produce_warnings(self, template_path: Path) -> None:
        result = read_template(template_path)
        for name in (SHEET_INSTRUCTIONS, SHEET_SAMPLE_REFERENCE):
            assert any(name in warning for warning in result.warnings)

    def test_excluded_sheets_do_not_appear_in_rows(self, template_path: Path) -> None:
        result = read_template(template_path)
        assert SHEET_INSTRUCTIONS not in result.rows
        assert SHEET_SAMPLE_REFERENCE not in result.rows
        assert "数据校验" not in result.rows

    def test_sample_sheet_data_is_not_imported(self, tmp_path: Path) -> None:
        """即使样本表被填了数据，也不得进入业务行。"""
        path = write_template(tmp_path / TEMPLATE_NAME)
        workbook = load_workbook(path)
        sheet = workbook[SHEET_SAMPLE_REFERENCE]
        sheet.cell(row=20, column=1, value="本期总电量_kWh")
        sheet.cell(row=20, column=2, value=8780498)
        # 同时给月度汇总写一条真实数据，确认两者不混淆
        _fill_row(workbook[SHEET_MONTHLY_SUMMARY], headers_of(SHEET_MONTHLY_SUMMARY), SEP_ROW)
        workbook.save(path)

        result = read_template(path)
        assert SHEET_SAMPLE_REFERENCE not in result.rows
        assert len(result.business_rows(SHEET_MONTHLY_SUMMARY)) == 1


# ---------------------------------------------------------------------------
# 39.8 重复识别所需信息
# ---------------------------------------------------------------------------


class TestDuplicateDetectionInputs:
    def test_file_hash_is_stable_and_content_sensitive(self, tmp_path: Path) -> None:
        """同一文件重复哈希必须一致；内容变化后哈希必须变化。

        注意：**不能比较两个分别生成的文件**——xlsx 是 zip 包，
        写入时间戳会使两次生成的字节不同（实测确认），
        因此"同一内容的两次生成"哈希并不相等。重复识别依赖的是
        **同一个文件**的哈希稳定性，而不是跨生成的确定性。
        """
        target = write_template(tmp_path / "a.xlsx")
        first_hash = compute_file_hash(target)
        assert compute_file_hash(target) == first_hash, "同一文件重复哈希必须一致"

        _append_monthly_row(target, SEP_ROW)
        assert compute_file_hash(target) != first_hash, "内容变化后哈希必须变化"

        stable_again = compute_file_hash(target)
        assert compute_file_hash(target) == stable_again

    def test_same_file_read_twice_yields_same_hash(self, template_path: Path) -> None:
        first = read_template(template_path)
        second = read_template(template_path)
        assert first.file_hash == second.file_hash
        assert first.file_name == second.file_name

    def test_duplicate_period_rows_are_both_read(self, template_path: Path) -> None:
        """重复账期由服务层判定；读取层必须如实读出两行，不得静默去重。"""
        _append_monthly_row(template_path, SEP_ROW, row=2)
        _append_monthly_row(template_path, SEP_ROW, row=3)
        rows = read_template(template_path).business_rows(SHEET_MONTHLY_SUMMARY)
        assert len(rows) == 2

    def test_row_carries_key_fields_for_duplicate_check(self, template_path: Path) -> None:
        _append_monthly_row(template_path, SEP_ROW)
        row = read_template(template_path).business_rows(SHEET_MONTHLY_SUMMARY)[0]
        # 项目 + 户号 + 账期 是重复判定的核心字段
        assert row.text("项目编号") == "PRJ-001"
        assert row.text("户号") == "1234567890"
        assert row.text("账期") == "2025-09"


# ---------------------------------------------------------------------------
# 39.9 有效/无效混合与结果序列化
# ---------------------------------------------------------------------------


class TestMixedValidityAndSerialization:
    def test_mixed_rows_are_all_read_for_later_validation(self, template_path: Path) -> None:
        _append_monthly_row(template_path, SEP_ROW, row=2)
        _append_monthly_row(template_path, {"项目编号": "PRJ-002", "账期": "2025-10"}, row=3)
        rows = read_template(template_path).business_rows(SHEET_MONTHLY_SUMMARY)
        assert len(rows) == 2

    def test_read_result_to_dict_is_serialisable(self, template_path: Path) -> None:
        _append_monthly_row(template_path, SEP_ROW)
        payload = read_template(template_path).to_dict()
        assert payload["file_name"] == TEMPLATE_NAME
        assert payload["template_version"] == "1.0"
        assert payload["row_counts"][SHEET_MONTHLY_SUMMARY] == 1
        assert isinstance(payload["warnings"], list)

    def test_row_to_dict_serialises_dates(self, template_path: Path) -> None:
        _append_monthly_row(template_path, SEP_ROW)
        row = read_template(template_path).business_rows(SHEET_MONTHLY_SUMMARY)[0]
        payload = row.to_dict()
        assert payload["sheet_name"] == SHEET_MONTHLY_SUMMARY
        assert payload["row_number"] == 2
        assert payload["cells"]["账期开始日期"] in ("2025-09-01", "2025-09-01T00:00:00")

    def test_max_rows_limit_is_respected(self, template_path: Path) -> None:
        for row in range(2, 12):
            _append_monthly_row(template_path, {**SEP_ROW, "账期": f"2025-{row:02d}"}, row=row)
        result = read_template(template_path, max_rows_per_sheet=3)
        assert len(result.business_rows(SHEET_MONTHLY_SUMMARY)) == 3

    def test_readonly_result_has_all_rows_helper(self, template_path: Path) -> None:
        _append_monthly_row(template_path, SEP_ROW)
        result = read_template(template_path)
        assert len(result.all_rows) == 1
        assert result.all_rows[0].sheet_name == SHEET_MONTHLY_SUMMARY

    def test_error_format_includes_details(self) -> None:
        error = TemplateReadError("表头不完整", details=["缺 A", "缺 B"])
        text = error.format()
        assert "表头不完整" in text
        assert "缺 A" in text
        assert "缺 B" in text
