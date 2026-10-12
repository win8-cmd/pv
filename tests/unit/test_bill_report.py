"""账单事实报告测试（规格书第 34、41 节）。

覆盖两类关注点：

**A. 报告规则（report_data）**

- 未提供的值显示"未提供"而**不是 0**；
- 合计只累加**已提供**的账单，缺失不参与求和；
- 费用基准小计只含"未包含在其他费用中"且"纳入基准电费=是"的项，
  无参与项时返回 ``None``（避免被读成"合计为零"）；
- 数据完整性逐项给出"满足 / 不满足 / 待核实"，**不做评分**；
- 数值一律 ``Decimal``，传入 ``float`` 被拒绝；
- 报告**不得**出现收益率/回收期等经济性结论。

**B. 文件产出（report_writer）**

- 6 张工作表齐备且顺序稳定；
- 金额/电量为**文本**写入（保留十进制精度，不经过浮点）；
- 无公式（本报告是事实快照，不含需要重算的单元格）；
- 通过 skill 的 Office 结构检查；
- 标题、来源、明细行数与数据一致。
"""

from __future__ import annotations

from datetime import date, time
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from cenep.data.report_data import (
    MISSING_TEXT,
    build_report_bundle,
    completeness_findings,
    fee_baseline_subtotal,
    format_optional,
    internal_fee_subtotal,
)
from cenep.data.report_writer import (
    REPORT_SHEET_NAMES,
    default_report_name,
    write_report,
)
from cenep.domain.bill_models import (
    Avoidability,
    BaselineInclusion,
    BillFeeItem,
    BillMeterPoint,
    BillTimeOfUseItem,
    ElectricityBill,
    FeeCategory,
    IncludedInOtherFee,
    MeterPointType,
    MeteringBasis,
    ReviewStatus,
    TimeOfUsePeriod,
    YesNoUnknown,
)
from cenep.domain.exceptions import PersistenceError


def _bill(**overrides) -> ElectricityBill:
    values = dict(
        bill_id="b1",
        project_id="p1",
        billing_period="2025-09",
        period_start=date(2025, 9, 1),
        period_end=date(2025, 9, 30),
        account_number="1234567890",
        total_energy_kwh=Decimal("8780498"),
        total_amount_yuan=Decimal("6699241.94"),
        power_factor_adjustment_fee_yuan=Decimal("-9206.37"),
        source_file="2025-09账单.pdf",
        source_page="1",
        review_status=ReviewStatus.REVIEWED,
    )
    values.update(overrides)
    return ElectricityBill(**values)


def _fee(**overrides) -> BillFeeItem:
    values = dict(
        bill_id="b1",
        project_id="p1",
        billing_period="2025-09",
        fee_name="输配电费",
        fee_category=FeeCategory.TRANSMISSION,
        amount_yuan=Decimal("2000662.04"),
        included_in_other=IncludedInOtherFee.NO,
        baseline_inclusion=BaselineInclusion.YES,
    )
    values.update(overrides)
    return BillFeeItem(**values)


def _bundle(bills, **kwargs):
    return build_report_bundle(
        project_id="p1",
        project_name="示例项目",
        bills=bills,
        generated_at="2026-10-12T00:00:00",
        **kwargs,
    )


# ---------------------------------------------------------------------------
# A. 报告规则
# ---------------------------------------------------------------------------


class TestMissingValueSemantics:
    def test_none_shows_missing_not_zero(self) -> None:
        assert format_optional(None) == MISSING_TEXT
        assert format_optional(None) != "0"

    def test_zero_shows_zero(self) -> None:
        assert format_optional(Decimal("0")) == "0"

    def test_float_is_rejected(self) -> None:
        """报告数值必须为 Decimal，防止浮点混入。"""
        with pytest.raises(TypeError, match="Decimal"):
            format_optional(6699241.94)  # type: ignore[arg-type]

    def test_totals_skip_missing_bills(self) -> None:
        """合计只累加已提供的账单，缺失不当作 0。"""
        bundle = _bundle(
            [
                _bill(total_amount_yuan=Decimal("100"), total_energy_kwh=Decimal("10")),
                _bill(
                    bill_id="b2",
                    account_number="2",
                    billing_period="2025-10",
                    total_amount_yuan=None,
                    total_energy_kwh=None,
                ),
            ]
        )
        assert bundle.total_amount_yuan == Decimal("100")
        assert bundle.total_energy_kwh == Decimal("10")

    def test_all_missing_totals_are_none_not_zero(self) -> None:
        bundle = _bundle([_bill(total_amount_yuan=None, total_energy_kwh=None)])
        assert bundle.total_amount_yuan is None
        assert bundle.total_energy_kwh is None
        assert format_optional(bundle.total_amount_yuan) == MISSING_TEXT


class TestFeeSemantics:
    def test_baseline_subtotal_only_counts_qualifying_items(self) -> None:
        items = [
            _fee(fee_name="输配电费", amount_yuan=Decimal("100")),
            _fee(
                fee_name="已包含项",
                amount_yuan=Decimal("999"),
                included_in_other=IncludedInOtherFee.YES,
                included_in_fee_name="输配电费",
            ),
            _fee(
                fee_name="不纳入基准",
                amount_yuan=Decimal("888"),
                baseline_inclusion=BaselineInclusion.NO,
            ),
        ]
        assert fee_baseline_subtotal(items) == Decimal("100")

    def test_baseline_subtotal_none_when_no_participants(self) -> None:
        """无参与项时返回 None，而不是 0（否则会被读成"合计为零"）。"""
        assert fee_baseline_subtotal([]) is None
        assert fee_baseline_subtotal([_fee(included_in_other=IncludedInOtherFee.TO_VERIFY)]) is None

    def test_internal_subtotal_lists_included_items(self) -> None:
        items = [
            _fee(
                fee_name="运营费",
                amount_yuan=Decimal("-141770"),
                included_in_other=IncludedInOtherFee.YES,
                included_in_fee_name="购电电费",
            )
        ]
        assert internal_fee_subtotal(items) == Decimal("-141770")

    def test_internal_subtotal_none_when_empty(self) -> None:
        assert internal_fee_subtotal([]) is None


class TestCompletenessFindings:
    def test_continuous_months_satisfied(self) -> None:
        bundle = _bundle(
            [
                _bill(billing_period="2025-09"),
                _bill(bill_id="b2", account_number="2", billing_period="2025-10"),
            ]
        )
        findings = dict((q, (v, d)) for q, v, d in completeness_findings(bundle))
        verdict, _detail = findings["是否覆盖连续月份（账期无缺月）"]
        assert verdict == "满足"

    def test_month_gap_is_not_satisfied(self) -> None:
        bundle = _bundle(
            [
                _bill(billing_period="2025-09"),
                _bill(bill_id="b2", account_number="2", billing_period="2025-12"),
            ]
        )
        findings = dict((q, (v, d)) for q, v, d in completeness_findings(bundle))
        verdict, detail = findings["是否覆盖连续月份（账期无缺月）"]
        assert verdict == "不满足"
        assert "2025-09→2025-12" in detail

    def test_year_boundary_is_continuous(self) -> None:
        bundle = _bundle(
            [
                _bill(billing_period="2025-12"),
                _bill(bill_id="b2", account_number="2", billing_period="2026-01"),
            ]
        )
        findings = dict((q, (v, d)) for q, v, d in completeness_findings(bundle))
        assert findings["是否覆盖连续月份（账期无缺月）"][0] == "满足"

    def test_unconfirmed_bills_not_satisfied(self) -> None:
        bundle = _bundle([_bill()])
        findings = dict((q, (v, d)) for q, v, d in completeness_findings(bundle))
        verdict, detail = findings["是否每张账单都已人工确认"]
        assert verdict == "不满足"
        assert "不得用于正式测算" in detail

    def test_confirmed_bills_satisfied(self) -> None:
        # confirm() 返回副本（不改原对象），必须接收返回值
        bill = _bill().confirm()
        bundle = _bundle([bill])
        findings = dict((q, (v, d)) for q, v, d in completeness_findings(bundle))
        assert findings["是否每张账单都已人工确认"][0] == "满足"

    def test_missing_source_not_satisfied(self) -> None:
        bundle = _bundle([_bill(source_page=None)])
        findings = dict((q, (v, d)) for q, v, d in completeness_findings(bundle))
        verdict, detail = findings["是否记录了来源文件与页码（可追溯）"]
        assert verdict == "不满足"
        assert "2025-09" in detail

    def test_no_tou_energy_is_not_satisfied(self) -> None:
        """缺少分时电量必须明确"无法评估消纳"，不得给出估算。"""
        bundle = _bundle([_bill()])
        findings = dict((q, (v, d)) for q, v, d in completeness_findings(bundle))
        verdict, detail = findings["是否有分时电量（缺少则无法评估光伏消纳）"]
        assert verdict == "不满足"
        assert "无法精确评估" in detail

    def test_with_tou_energy_satisfied(self) -> None:
        item = BillTimeOfUseItem(
            bill_id="b1", project_id="p1", billing_period="2025-09",
            metering_basis_id="MP1", basis=MeteringBasis.PURCHASE,
            period_name=TimeOfUsePeriod.FLAT, energy_kwh=Decimal("100"),
        )
        bundle = _bundle([_bill()], tou_items=[item])
        findings = dict((q, (v, d)) for q, v, d in completeness_findings(bundle))
        assert findings["是否有分时电量（缺少则无法评估光伏消纳）"][0] == "满足"

    def test_all_five_questions_are_answered(self) -> None:
        bundle = _bundle([_bill()])
        assert len(completeness_findings(bundle)) == 5

    def test_no_scoring_only_verdicts(self) -> None:
        """只允许三种结论，不得出现分数或百分比。"""
        bundle = _bundle([_bill(), _bill(bill_id="b2", account_number="2", billing_period="2025-10")])
        for _question, verdict, _detail in completeness_findings(bundle):
            assert verdict in ("满足", "不满足", "待核实")


class TestNoEconomicConclusions:
    def test_notes_forbid_average_price_as_savings(self) -> None:
        """必须显式声明账单平均电价不能当作光伏边际节省电价。"""
        bundle = _bundle([_bill()])
        joined = "".join(bundle.notes)
        assert "不能作为光伏边际节省电价" in joined

    def test_notes_forbid_economic_conclusions(self) -> None:
        bundle = _bundle([_bill()])
        joined = "".join(bundle.notes)
        assert "不包含收益率" in joined

    def test_bundle_dict_has_no_economic_metrics(self) -> None:
        payload = _bundle([_bill()]).to_dict()
        for forbidden in ("irr", "npv", "payback", "roi", "收益率", "回收期"):
            assert forbidden not in " ".join(payload).lower()

    def test_average_price_is_reported_as_fact_only(self) -> None:
        """平均电价在账单事实表里作为原始字段出现，且不参与任何计算列。"""
        from cenep.data.report_data import BILL_FACT_COLUMNS

        titles = [title for title, _ in BILL_FACT_COLUMNS]
        assert "账单展示平均电价_元每kWh" in titles
        # 不得存在"节省电价""边际电价"之类推导列
        assert not any("节省" in title or "边际" in title for title in titles)


class TestBundleBasics:
    def test_counts_and_periods(self) -> None:
        second = _bill(
            bill_id="b2", billing_period="2025-10", account_number="2"
        ).confirm()
        bundle = _bundle([_bill(), second])
        assert bundle.bill_count == 2
        assert bundle.confirmed_count == 1
        assert bundle.usable_count == 1
        assert bundle.periods == ("2025-09", "2025-10")

    def test_validation_is_computed(self) -> None:
        bundle = _bundle([_bill()])
        assert bundle.validation is not None
        assert bundle.validation.checked_bills == 1

    def test_sample_bill_not_usable(self) -> None:
        """样本账单即使试图确认也不可用于测算。"""
        bill = _bill(bill_id="sample-1", is_sample=True)
        with pytest.raises(ValueError, match="样本"):
            bill.confirm()
        bundle = _bundle([bill])
        assert bundle.usable_count == 0


# ---------------------------------------------------------------------------
# C. 端到端：应用服务导出（SQLite -> 报告文件）
# ---------------------------------------------------------------------------


class TestServiceExportEndToEnd:
    def _service(self, tmp_path):
        from cenep.application.bill_service import BillService
        from cenep.domain.models import Project
        from cenep.infrastructure.database import Database, ProjectRepository

        db = Database(tmp_path / "svc.db")
        db.initialize_schema()
        project = ProjectRepository(db).add(Project(name="导出测试项目"))
        return BillService(db), db, project.project_id

    def test_export_reads_from_database(self, tmp_path) -> None:
        from cenep.application.bill_service import ManualBillInput

        service, db, project_id = self._service(tmp_path)
        service.create_bill(
            project_id,
            ManualBillInput(
                billing_period="2025-09",
                period_start="2025-09-01",
                period_end="2025-09-30",
                account_number="1234567890",
                total_energy_kwh="8780498",
                total_amount_yuan="6699241.94",
                power_factor_adjustment_fee_yuan="-9206.37",
                source_file="a.pdf",
                source_page="1",
                review_status="已复核",
            ),
        )
        path = service.export_report(
            project_id,
            tmp_path / "out.xlsx",
            project_name="导出测试项目",
            generated_at="2026-10-12T00:00:00",
        )
        assert path.exists()
        workbook = load_workbook(path)
        sheet = workbook["账单事实"]
        titles = [cell.value for cell in sheet[1]]
        column = titles.index("账单本期电费_元") + 1
        assert sheet.cell(row=2, column=column).value == "6699241.94"
        neg = titles.index("功率因数调整电费_元") + 1
        assert sheet.cell(row=2, column=neg).value == "-9206.37"
        db.close()

    def test_export_empty_project_still_produces_workbook(self, tmp_path) -> None:
        """无账单也应产出结构完整的报告（界面层负责拦截"无意义导出"）。"""
        service, db, project_id = self._service(tmp_path)
        path = service.export_report(project_id, tmp_path / "empty.xlsx")
        workbook = load_workbook(path)
        assert tuple(workbook.sheetnames) == REPORT_SHEET_NAMES
        assert workbook["账单事实"].max_row == 1, "无数据时只有表头"
        db.close()

    def test_build_report_does_not_write_file(self, tmp_path) -> None:
        from cenep.application.bill_service import ManualBillInput

        service, db, project_id = self._service(tmp_path)
        service.create_bill(
            project_id,
            ManualBillInput(
                billing_period="2025-09",
                account_number="1",
                total_energy_kwh="1",
                total_amount_yuan="2",
            ),
        )
        bundle = service.build_report(project_id, project_name="X")
        assert bundle.project_name == "X"
        assert list(tmp_path.glob("*.xlsx")) == [], "build_report 不得写文件"
        db.close()


# ---------------------------------------------------------------------------
# B. 文件产出
# ---------------------------------------------------------------------------


class TestWorkbookStructure:
    def test_writes_six_sheets_in_order(self, tmp_path) -> None:
        path = write_report(_bundle([_bill()]), tmp_path / "r.xlsx")
        workbook = load_workbook(path)
        assert tuple(workbook.sheetnames) == REPORT_SHEET_NAMES

    def test_extension_is_corrected(self, tmp_path) -> None:
        path = write_report(_bundle([_bill()]), tmp_path / "r.txt")
        assert path.suffix == ".xlsx"
        assert path.exists()

    def test_parent_directory_is_created(self, tmp_path) -> None:
        path = write_report(_bundle([_bill()]), tmp_path / "a" / "b" / "r.xlsx")
        assert path.exists()

    def test_no_macros_in_package(self, tmp_path) -> None:
        import zipfile

        path = write_report(_bundle([_bill()]), tmp_path / "r.xlsx")
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
        assert not any("vbaProject" in name for name in names)
        assert not any(name.endswith(".bin") for name in names)

    def test_no_formulas(self, tmp_path) -> None:
        """报告是事实快照，不应含需要重算的公式。"""
        path = write_report(_bundle([_bill()]), tmp_path / "r.xlsx")
        workbook = load_workbook(path)
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    assert cell.data_type != "f", f"{sheet.title}!{cell.coordinate} 含公式"

    def test_default_report_name_is_sanitised(self) -> None:
        name = default_report_name("工厂/A:B*C?", generated_on="2026-10-12")
        assert "/" not in name and ":" not in name and "*" not in name and "?" not in name
        assert name.endswith(".xlsx")


class TestWorkbookContent:
    def test_missing_value_appears_as_text_not_zero(self, tmp_path) -> None:
        """未提供的金额必须在表里显示"未提供"，不得出现 0。"""
        path = write_report(
            _bundle([_bill(vat_invoice_amount_yuan=None)]), tmp_path / "r.xlsx"
        )
        workbook = load_workbook(path)
        sheet = workbook["账单事实"]
        titles = [cell.value for cell in sheet[1]]
        column = titles.index("增值税专票金额_元") + 1
        assert sheet.cell(row=2, column=column).value == MISSING_TEXT

    def test_amounts_are_text_preserving_precision(self, tmp_path) -> None:
        """金额以文本写入，避免浮点精度损失。"""
        path = write_report(
            _bundle([_bill(total_amount_yuan=Decimal("6699241.94"))]), tmp_path / "r.xlsx"
        )
        workbook = load_workbook(path)
        sheet = workbook["账单事实"]
        titles = [cell.value for cell in sheet[1]]
        column = titles.index("账单本期电费_元") + 1
        cell = sheet.cell(row=2, column=column)
        assert cell.value == "6699241.94"
        assert isinstance(cell.value, str), "金额应为文本，避免浮点误差"

    def test_negative_fee_preserved_in_report(self, tmp_path) -> None:
        path = write_report(_bundle([_bill()]), tmp_path / "r.xlsx")
        workbook = load_workbook(path)
        sheet = workbook["账单事实"]
        titles = [cell.value for cell in sheet[1]]
        column = titles.index("功率因数调整电费_元") + 1
        assert sheet.cell(row=2, column=column).value == "-9206.37"

    def test_bill_rows_match_data(self, tmp_path) -> None:
        bills = [
            _bill(),
            _bill(bill_id="b2", account_number="2", billing_period="2025-10"),
        ]
        path = write_report(_bundle(bills), tmp_path / "r.xlsx")
        workbook = load_workbook(path)
        sheet = workbook["账单事实"]
        # 表头 1 行 + 2 条数据
        assert sheet.max_row == 3

    def test_completeness_sheet_has_five_rows(self, tmp_path) -> None:
        path = write_report(_bundle([_bill()]), tmp_path / "r.xlsx")
        workbook = load_workbook(path)
        sheet = workbook["数据完整性"]
        assert sheet.max_row == 6  # 表头 + 5 项

    def test_notes_sheet_states_boundaries(self, tmp_path) -> None:
        path = write_report(_bundle([_bill()]), tmp_path / "r.xlsx")
        workbook = load_workbook(path)
        sheet = workbook["报告说明"]
        texts = [
            str(cell.value)
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        ]
        joined = " ".join(texts)
        assert "不能作为光伏边际节省电价" in joined
        assert MISSING_TEXT in joined

    def test_overview_reports_zero_usable_when_unconfirmed(self, tmp_path) -> None:
        path = write_report(_bundle([_bill()]), tmp_path / "r.xlsx")
        workbook = load_workbook(path)
        sheet = workbook["数据概览"]
        rows = {str(row[0]): row[1] for row in sheet.iter_rows(values_only=True) if row and row[0]}
        assert rows["可用于正式测算"] == "0 / 1"

    def test_fee_sheet_marks_baseline_inclusion(self, tmp_path) -> None:
        fees = [
            _fee(fee_name="输配电费", amount_yuan=Decimal("100")),
            _fee(
                fee_name="运营费",
                amount_yuan=Decimal("50"),
                included_in_other=IncludedInOtherFee.YES,
                included_in_fee_name="输配电费",
            ),
        ]
        path = write_report(_bundle([_bill()], fee_items=fees), tmp_path / "r.xlsx")
        workbook = load_workbook(path)
        sheet = workbook["费用明细"]
        titles = [cell.value for cell in sheet[1]]
        column = titles.index("是否计入基准合计") + 1
        flags = [sheet.cell(row=row, column=column).value for row in (2, 3)]
        assert flags == ["是", "否"], "已包含在其他费用中的项不得计入基准合计"

    def test_validation_detail_present(self, tmp_path) -> None:
        path = write_report(_bundle([_bill()]), tmp_path / "r.xlsx")
        workbook = load_workbook(path)
        sheet = workbook["分时与计量点"]
        texts = [
            str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value
        ]
        joined = " ".join(texts)
        assert "校验明细" in joined
        assert "V01" in joined, "13 类校验明细应写入报告"

    def test_tou_and_meter_rows_written(self, tmp_path) -> None:
        tou = BillTimeOfUseItem(
            bill_id="b1", project_id="p1", billing_period="2025-09",
            metering_basis_id="MP1", basis=MeteringBasis.PURCHASE,
            period_name=TimeOfUsePeriod.FLAT,
            segment_start=time(8, 0), segment_end=time(22, 0),
            energy_kwh=Decimal("100.5"), price_yuan_per_kwh=Decimal("0.76297"),
        )
        meter = BillMeterPoint(
            bill_id="b1", project_id="p1", billing_period="2025-09",
            point_code="M1", point_type=MeterPointType.SUB,
            energy_kwh=Decimal("8516838"),
            included_in_total=YesNoUnknown.NO, allowed_in_rollup=YesNoUnknown.YES,
        )
        path = write_report(
            _bundle([_bill()], tou_items=[tou], meter_points=[meter]), tmp_path / "r.xlsx"
        )
        workbook = load_workbook(path)
        sheet = workbook["分时与计量点"]
        texts = [
            str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value
        ]
        joined = " ".join(texts)
        assert "100.5" in joined
        assert "M1" in joined
        assert "0.76297" in joined
