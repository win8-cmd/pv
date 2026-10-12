"""`bill-template` 命令测试（规格书第 12 节"下载模板"）。

命令行是模板最直接的获取入口，且不依赖界面。
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from cenep.cli import main
from cenep.data.bill_template_spec import TEMPLATE_NAME


class TestBillTemplateCommand:
    def test_command_generates_template_in_directory(
        self, tmp_path: Path, capsys
    ) -> None:
        exit_code = main(["bill-template", "--dir", str(tmp_path)])
        assert exit_code == 0
        output = capsys.readouterr().out
        assert "模板已生成" in output
        assert "样本参考" in output

        generated = tmp_path / TEMPLATE_NAME
        assert generated.exists()
        workbook = load_workbook(generated)
        assert workbook.sheetnames[0] == "使用说明"
        assert len(workbook.sheetnames) == 7

    def test_command_honours_explicit_output_path(
        self, tmp_path: Path, capsys
    ) -> None:
        target = tmp_path / "custom_name.xlsx"
        exit_code = main(["bill-template", "--out", str(target)])
        assert exit_code == 0
        assert target.exists()
        assert "custom_name.xlsx" in capsys.readouterr().out

    def test_explicit_output_wins_over_directory(self, tmp_path: Path) -> None:
        """同时给出 --out 与 --dir 时以 --out 为准。"""
        target = tmp_path / "explicit.xlsx"
        other_dir = tmp_path / "other"
        exit_code = main(
            ["bill-template", "--out", str(target), "--dir", str(other_dir)]
        )
        assert exit_code == 0
        assert target.exists()
        assert not other_dir.exists(), "指定 --out 时不应再往 --dir 写文件"

    def test_command_creates_missing_parent_directory(self, tmp_path: Path) -> None:
        target = tmp_path / "a" / "b" / TEMPLATE_NAME
        assert main(["bill-template", "--out", str(target)]) == 0
        assert target.exists()

    def test_generated_template_is_readable_and_macro_free(self, tmp_path: Path) -> None:
        main(["bill-template", "--dir", str(tmp_path)])
        generated = tmp_path / TEMPLATE_NAME
        assert generated.suffix == ".xlsx"
        workbook = load_workbook(generated)
        assert "月度汇总" in workbook.sheetnames
        # 业务表只应有表头行
        sheet = workbook["月度汇总"]
        assert sheet.max_row == 1
