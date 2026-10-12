"""账单 SQLite 持久化测试（规格书第 34 节"保存后重开仍可恢复"）。

覆盖：

- 六类实体写入后**关闭数据库再重开**仍可完整读回；
- **金额与电量以 ``Decimal`` 往返不丢精度**（不经 float）；
- **``None`` 与 ``Decimal(0)`` 往返后仍严格区分**；
- **负费用完整保留**；
- 枚举、日期、时间往返一致；
- 导入服务 + SQLite 仓储的端到端链路；
- 重复导入的跳过 / 替换 / 保留在**真实数据库**上的行为；
- 级联删除与项目删除保护；
- 人工确认与样本不得确认；
- 旧库（v1）可打开并自动迁移（在模式版本测试中另有覆盖）。
"""

from __future__ import annotations

from datetime import date, time
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from cenep.data.bill_import_service import (
    BillImportService,
    DuplicateDecision,
    ImportDecision,
)
from cenep.data.bill_template_spec import (
    SHEET_FEE_ITEMS,
    SHEET_METER_POINTS,
    SHEET_MONTHLY_SUMMARY,
    SHEET_TIME_OF_USE,
    TEMPLATE_NAME,
    headers_of,
)
from cenep.data.template_writer import write_template
from cenep.domain.bill_models import (
    Avoidability,
    BaselineInclusion,
    BillFeeItem,
    BillImportBatch,
    BillMeterPoint,
    BillTimeOfUseItem,
    ElectricityBill,
    FeeCategory,
    ImportOutcome,
    IncludedInOtherFee,
    MeterPointType,
    MeteringBasis,
    ReviewStatus,
    TimeOfUsePeriod,
    YesNoUnknown,
)
from cenep.domain.models import Project
from cenep.infrastructure.bill_repository import SqliteBillRepository
from cenep.infrastructure.database import Database, ProjectRepository

#: 项目主键由模型生成：校验规则要求 32 位十六进制，不能用字面量
_PROJECT = Project(name="持久化测试项目")
PROJECT = _PROJECT.project_id


@pytest.fixture
def db_path(tmp_path: Path) -> Path:
    return tmp_path / "cenep.db"


@pytest.fixture
def database(db_path: Path):
    db = Database(db_path)
    db.initialize_schema()
    ProjectRepository(db).add(_PROJECT)
    yield db
    db.close()


def _full_bill(**overrides) -> ElectricityBill:
    values = dict(
        bill_id="bill-1",
        project_id=PROJECT,
        customer_name="示例用户",
        account_number="1234567890",
        billing_period="2025-09",
        period_start=date(2025, 9, 1),
        period_end=date(2025, 9, 30),
        total_energy_kwh=Decimal("8780498"),
        total_amount_yuan=Decimal("6699241.94"),
        displayed_average_price=Decimal("0.76297"),
        power_factor_adjustment_fee_yuan=Decimal("-9206.37"),
        demand_charge_yuan=Decimal("1202916"),
        basic_charge_yuan=Decimal("0"),
        actual_demand=Decimal("30844"),
        actual_demand_unit="kW",
        source_file="2025-09账单.pdf",
        source_page="1",
        review_status=ReviewStatus.REVIEWED,
    )
    values.update(overrides)
    return ElectricityBill(**values)


# ---------------------------------------------------------------------------
# 基本往返
# ---------------------------------------------------------------------------


class TestBillRoundTrip:
    def test_bill_survives_reopen(self, db_path: Path) -> None:
        db = Database(db_path)
        db.initialize_schema()
        ProjectRepository(db).add(_PROJECT)
        repository = SqliteBillRepository(db)
        repository.save_batch(
            BillImportBatch(project_id=PROJECT, file_name="a.xlsx", file_hash="hash-1"),
            [_full_bill()],
            [],
            [],
            [],
        )
        db.close()

        # 重新打开（模拟"保存后重开软件"）
        reopened = Database(db_path)
        reopened.initialize_schema()
        repository = SqliteBillRepository(reopened)
        bills = repository.list_bills(PROJECT)
        assert len(bills) == 1
        bill = bills[0]
        assert bill.bill_id == "bill-1"
        assert bill.billing_period == "2025-09"
        assert bill.period_start == date(2025, 9, 1)
        assert bill.period_end == date(2025, 9, 30)
        assert bill.review_status is ReviewStatus.REVIEWED
        assert bill.actual_demand_unit == "kW"
        reopened.close()

    def test_amount_round_trips_as_decimal_exactly(self, db_path: Path) -> None:
        """金额必须精确往返，不得经过 float。"""
        db = Database(db_path)
        db.initialize_schema()
        ProjectRepository(db).add(_PROJECT)
        repository = SqliteBillRepository(db)
        repository.save_batch(
            BillImportBatch(project_id=PROJECT, file_name="a.xlsx"),
            [_full_bill(total_amount_yuan=Decimal("6699241.94"))],
            [], [], [],
        )
        db.close()

        reopened = Database(db_path)
        reopened.initialize_schema()
        bill = SqliteBillRepository(reopened).get_bill("bill-1")
        assert bill is not None
        assert isinstance(bill.total_amount_yuan, Decimal)
        assert bill.total_amount_yuan == Decimal("6699241.94")
        assert str(bill.total_amount_yuan) == "6699241.94", "不得出现浮点尾数"
        reopened.close()

    def test_negative_fee_round_trips(self, db_path: Path) -> None:
        db = Database(db_path)
        db.initialize_schema()
        ProjectRepository(db).add(_PROJECT)
        SqliteBillRepository(db).save_batch(
            BillImportBatch(project_id=PROJECT, file_name="a.xlsx"),
            [_full_bill()], [], [], [],
        )
        db.close()

        reopened = Database(db_path)
        reopened.initialize_schema()
        bill = SqliteBillRepository(reopened).get_bill("bill-1")
        assert bill.power_factor_adjustment_fee_yuan == Decimal("-9206.37")
        assert bill.power_factor_adjustment_fee_yuan < 0
        reopened.close()

    def test_none_and_zero_are_distinguished(self, db_path: Path) -> None:
        """``None``（未提供）与 ``Decimal(0)``（确为零）必须严格区分。"""
        db = Database(db_path)
        db.initialize_schema()
        ProjectRepository(db).add(_PROJECT)
        repository = SqliteBillRepository(db)
        repository.save_batch(
            BillImportBatch(project_id=PROJECT, file_name="a.xlsx"),
            [
                _full_bill(
                    bill_id="bill-none",
                    account_number="acc-none",
                    vat_invoice_amount_yuan=None,
                    basic_charge_yuan=Decimal("0"),
                )
            ],
            [], [], [],
        )
        db.close()

        reopened = Database(db_path)
        reopened.initialize_schema()
        bill = SqliteBillRepository(reopened).get_bill("bill-none")
        assert bill.vat_invoice_amount_yuan is None, "未提供必须读回 None"
        assert bill.vat_invoice_amount_yuan != Decimal(0)
        assert bill.basic_charge_yuan == Decimal(0), "确为零必须读回 Decimal(0)"
        assert bill.basic_charge_yuan is not None
        reopened.close()

    def test_batch_round_trips(self, db_path: Path) -> None:
        db = Database(db_path)
        db.initialize_schema()
        ProjectRepository(db).add(_PROJECT)
        repository = SqliteBillRepository(db)
        batch = BillImportBatch(
            project_id=PROJECT,
            file_name="2025-09.xlsx",
            file_hash="abc123",
            template_version="1.0",
            imported_by="tester",
            total_records=3,
            valid_records=2,
            warning_records=1,
            invalid_records=0,
            skipped_records=1,
            outcome=ImportOutcome.PARTIAL,
            user_decision="仅导入有效记录",
        )
        repository.save_batch(batch, [_full_bill()], [], [], [])
        db.close()

        reopened = Database(db_path)
        reopened.initialize_schema()
        batches = SqliteBillRepository(reopened).list_batches(PROJECT)
        assert len(batches) == 1
        stored = batches[0]
        assert stored.file_name == "2025-09.xlsx"
        assert stored.file_hash == "abc123"
        assert stored.outcome is ImportOutcome.PARTIAL
        assert stored.user_decision == "仅导入有效记录"
        assert stored.skipped_records == 1
        reopened.close()


# ---------------------------------------------------------------------------
# 明细表往返
# ---------------------------------------------------------------------------


class TestDetailRoundTrip:
    def _save_with_details(self, database: Database) -> None:
        repository = SqliteBillRepository(database)
        repository.save_batch(
            BillImportBatch(project_id=PROJECT, file_name="a.xlsx"),
            [_full_bill()],
            [
                BillTimeOfUseItem(
                    bill_id="bill-1",
                    project_id=PROJECT,
                    billing_period="2025-09",
                    metering_basis_id="MP1",
                    basis=MeteringBasis.PURCHASE,
                    period_name=TimeOfUsePeriod.FLAT,
                    segment_start=time(8, 0),
                    segment_end=time(9, 0),
                    energy_kwh=Decimal("100.5"),
                    price_yuan_per_kwh=Decimal("0.76297"),
                )
            ],
            [
                BillFeeItem(
                    bill_id="bill-1",
                    project_id=PROJECT,
                    billing_period="2025-09",
                    fee_category=FeeCategory.TRANSMISSION,
                    fee_name="输配电费",
                    amount_yuan=Decimal("2000662.04"),
                    included_in_other=IncludedInOtherFee.NO,
                    baseline_inclusion=BaselineInclusion.YES,
                )
            ],
            [
                BillMeterPoint(
                    bill_id="bill-1",
                    project_id=PROJECT,
                    billing_period="2025-09",
                    point_code="M1",
                    point_type=MeterPointType.SUB,
                    energy_kwh=Decimal("8516838"),
                    included_in_total=YesNoUnknown.NO,
                    allowed_in_rollup=YesNoUnknown.YES,
                )
            ],
        )

    def test_tou_item_round_trips(self, database: Database) -> None:
        self._save_with_details(database)
        items = SqliteBillRepository(database).list_tou_items("bill-1")
        assert len(items) == 1
        item = items[0]
        assert item.basis is MeteringBasis.PURCHASE
        assert item.period_name is TimeOfUsePeriod.FLAT
        assert item.segment_start == time(8, 0)
        assert item.segment_end == time(9, 0)
        assert item.energy_kwh == Decimal("100.5")
        assert item.price_yuan_per_kwh == Decimal("0.76297")

    def test_fee_item_round_trips(self, database: Database) -> None:
        self._save_with_details(database)
        items = SqliteBillRepository(database).list_fee_items("bill-1")
        assert len(items) == 1
        item = items[0]
        assert item.fee_category is FeeCategory.TRANSMISSION
        assert item.fee_name == "输配电费"
        assert item.amount_yuan == Decimal("2000662.04")
        assert item.included_in_other is IncludedInOtherFee.NO
        assert item.counts_toward_baseline is True, "落库后仍应可判定是否计入基准"

    def test_meter_point_round_trips(self, database: Database) -> None:
        self._save_with_details(database)
        items = SqliteBillRepository(database).list_meter_points("bill-1")
        assert len(items) == 1
        item = items[0]
        assert item.point_code == "M1"
        assert item.point_type is MeterPointType.SUB
        assert item.energy_kwh == Decimal("8516838")
        assert item.can_join_rollup is True


# ---------------------------------------------------------------------------
# 端到端：Excel → 导入服务 → SQLite → 重开
# ---------------------------------------------------------------------------


def _fill(path: Path, sheet: str, rows: list[dict], start: int = 2) -> None:
    workbook = load_workbook(path)
    headers = headers_of(sheet)
    for offset, values in enumerate(rows):
        excel_row = start + offset
        for index, header in enumerate(headers, start=1):
            if header in values:
                workbook[sheet].cell(row=excel_row, column=index, value=values[header])
    workbook.save(path)


SEP_ROW = {
    "项目编号": "PRJ-001", "户号": "1234567890", "账期": "2025-09",
    "账期开始日期": "2025-09-01", "账期结束日期": "2025-09-30",
    "本期总电量_kWh": 8780498, "账单本期电费_元": 6699241.94,
    "功率因数调整电费_元": -9206.37, "增值税专票金额_元": 6534168.58,
    "来源文件": "2025-09账单.pdf", "来源页码": "1", "复核状态": "已复核",
}
OCT_ROW = {
    **SEP_ROW, "账期": "2025-10", "账期开始日期": "2025-10-01",
    "账期结束日期": "2025-10-31", "本期总电量_kWh": 7479418,
    "账单本期电费_元": 5471334.24, "来源文件": "2025-10账单.pdf",
}


class TestEndToEndPersistence:
    def test_two_month_sample_regression(self, tmp_path: Path, db_path: Path) -> None:
        """两个月样本的端到端回归：导入 → 落库 → 重开仍可读回。"""
        template = write_template(tmp_path / TEMPLATE_NAME)
        _fill(template, SHEET_MONTHLY_SUMMARY, [SEP_ROW, OCT_ROW])
        _fill(
            template,
            SHEET_TIME_OF_USE,
            [
                {"项目编号": "PRJ-001", "账期": "2025-09", "计量口径/计量点编号": "MP1",
                 "电量类型": "购电", "时段名称": "平", "电量_kWh": 8780498},
                {"项目编号": "PRJ-001", "账期": "2025-10", "计量口径/计量点编号": "MP1",
                 "电量类型": "购电", "时段名称": "平", "电量_kWh": 7479418},
            ],
        )
        _fill(
            template,
            SHEET_FEE_ITEMS,
            [
                {"项目编号": "PRJ-001", "账期": "2025-09", "费用项目名称": "输配电费",
                 "金额_元": 2000662.04, "是否包含在其他费用中": "否",
                 "是否纳入基准电费": "是"},
                {"项目编号": "PRJ-001", "账期": "2025-09", "费用项目名称": "市场化运营费",
                 "金额_元": -141770, "是否包含在其他费用中": "是",
                 "被包含于/关联费用": "输配电费"},
            ],
        )
        _fill(
            template,
            SHEET_METER_POINTS,
            [
                {"项目编号": "PRJ-001", "账期": "2025-09", "计量点编号": "M1",
                 "计量点类型": "分表", "电量_kWh": 8516838,
                 "是否已计入首页总电量": "否", "是否允许纳入汇总": "是"},
            ],
        )

        db = Database(db_path)
        db.initialize_schema()
        ProjectRepository(db).add(_PROJECT)
        service = BillImportService(SqliteBillRepository(db))

        preview = service.preview(template, project_id=PROJECT)
        assert preview.counts.total == 2
        result = service.apply(preview, decision=ImportDecision.VALID_ONLY)
        assert result.imported_count == 2
        db.close()

        # 重开并核对
        reopened = Database(db_path)
        reopened.initialize_schema()
        repository = SqliteBillRepository(reopened)
        bills = repository.list_bills(PROJECT)
        assert [b.billing_period for b in bills] == ["2025-09", "2025-10"]

        september = bills[0]
        assert september.total_energy_kwh == Decimal("8780498")
        assert september.total_amount_yuan == Decimal("6699241.94")
        assert september.vat_invoice_amount_yuan == Decimal("6534168.58")
        assert september.power_factor_adjustment_fee_yuan == Decimal("-9206.37")
        assert september.review_status is ReviewStatus.REVIEWED
        # 导入的账单必须仍是"待确认"，未确认不得用于正式测算
        assert september.is_usable_for_calculation is False

        tou = repository.list_tou_items(september.bill_id)
        assert len(tou) == 1 and tou[0].energy_kwh == Decimal("8780498")
        fees = repository.list_fee_items(september.bill_id)
        assert len(fees) == 2
        assert any(f.amount_yuan == Decimal("-141770") for f in fees), "负费用保留"
        # 已包含在其他费用中的项不得计入基准合计
        included = [f for f in fees if f.included_in_other is IncludedInOtherFee.YES]
        assert included and included[0].counts_toward_baseline is False
        points = repository.list_meter_points(september.bill_id)
        assert len(points) == 1 and points[0].can_join_rollup is True
        reopened.close()

    def test_reimport_same_file_skip_keeps_one_bill(self, tmp_path: Path, db_path: Path) -> None:
        template = write_template(tmp_path / TEMPLATE_NAME)
        _fill(template, SHEET_MONTHLY_SUMMARY, [SEP_ROW])

        db = Database(db_path)
        db.initialize_schema()
        ProjectRepository(db).add(_PROJECT)
        service = BillImportService(SqliteBillRepository(db))

        first = service.preview(template, project_id=PROJECT)
        service.apply(first, decision=ImportDecision.VALID_ONLY)
        assert SqliteBillRepository(db).count_bills(PROJECT) == 1

        second = service.preview(template, project_id=PROJECT)
        assert second.has_duplicates, "重复导入必须被识别"
        service.apply(
            second,
            decision=ImportDecision.VALID_ONLY,
            duplicate_decision=DuplicateDecision.SKIP,
        )
        assert SqliteBillRepository(db).count_bills(PROJECT) == 1, "跳过不得新增记录"
        db.close()

    def test_reimport_replace_keeps_one_bill(self, tmp_path: Path, db_path: Path) -> None:
        template = write_template(tmp_path / TEMPLATE_NAME)
        _fill(template, SHEET_MONTHLY_SUMMARY, [SEP_ROW])

        db = Database(db_path)
        db.initialize_schema()
        ProjectRepository(db).add(_PROJECT)
        repository = SqliteBillRepository(db)
        service = BillImportService(repository)

        service.apply(
            service.preview(template, project_id=PROJECT), decision=ImportDecision.VALID_ONLY
        )
        second = service.preview(template, project_id=PROJECT)
        service.apply(
            second,
            decision=ImportDecision.VALID_ONLY,
            duplicate_decision=DuplicateDecision.REPLACE,
        )
        assert repository.count_bills(PROJECT) == 1, "替换后同一业务键只保留一条"
        db.close()

    def test_reimport_keep_adds_second_bill_is_blocked_by_unique_index(
        self, tmp_path: Path, db_path: Path
    ) -> None:
        """业务键唯一索引会拒绝"保留"造成的重复行。

        "保留"在导入服务层被允许（携带用户决策），但数据库层用**唯一索引**
        兜底：同一 项目+户号+账期 不允许两条。这保证即使调用方误用，
        数据库中也不会出现重复账期数据。
        """
        template = write_template(tmp_path / TEMPLATE_NAME)
        _fill(template, SHEET_MONTHLY_SUMMARY, [SEP_ROW])

        db = Database(db_path)
        db.initialize_schema()
        ProjectRepository(db).add(_PROJECT)
        service = BillImportService(SqliteBillRepository(db))
        service.apply(
            service.preview(template, project_id=PROJECT), decision=ImportDecision.VALID_ONLY
        )
        second = service.preview(template, project_id=PROJECT)

        from cenep.domain.exceptions import PersistenceError

        with pytest.raises(PersistenceError):
            service.apply(
                second,
                decision=ImportDecision.VALID_ONLY,
                duplicate_decision=DuplicateDecision.KEEP,
            )
        assert SqliteBillRepository(db).count_bills(PROJECT) == 1
        db.close()


# ---------------------------------------------------------------------------
# 安全与状态
# ---------------------------------------------------------------------------


class TestSafetyAndStatus:
    def test_project_delete_is_restricted_by_bills(self, database: Database) -> None:
        """存在账单时应拒绝删除项目（延续 RESTRICT 安全策略）。"""
        from cenep.domain.exceptions import PersistenceError

        SqliteBillRepository(database).save_batch(
            BillImportBatch(project_id=PROJECT, file_name="a.xlsx"),
            [_full_bill()], [], [], [],
        )
        with pytest.raises(PersistenceError):
            ProjectRepository(database).delete(PROJECT)

    def test_confirm_bill_sets_confirmed_state(self, database: Database) -> None:
        repository = SqliteBillRepository(database)
        repository.save_batch(
            BillImportBatch(project_id=PROJECT, file_name="a.xlsx"),
            [_full_bill()], [], [], [],
        )
        assert repository.get_bill("bill-1").confirmation.value == "待确认"
        repository.confirm_bill("bill-1", confirmed_by="tester")
        confirmed = repository.get_bill("bill-1")
        assert confirmed.confirmation.value == "已确认"

    def test_sample_bill_cannot_be_confirmed(self, database: Database) -> None:
        """样本数据不得确认（规格书第 30 节）。"""
        repository = SqliteBillRepository(database)
        repository.save_batch(
            BillImportBatch(project_id=PROJECT, file_name="sample.xlsx"),
            [_full_bill(is_sample=True)], [], [], [],
        )
        with pytest.raises(ValueError, match="样本"):
            repository.confirm_bill("bill-1")

    def test_update_review_status(self, database: Database) -> None:
        repository = SqliteBillRepository(database)
        repository.save_batch(
            BillImportBatch(project_id=PROJECT, file_name="a.xlsx"),
            [_full_bill()], [], [], [],
        )
        repository.update_review_status("bill-1", ReviewStatus.QUESTIONABLE)
        assert repository.get_bill("bill-1").review_status is ReviewStatus.QUESTIONABLE

    def test_confirm_missing_bill_raises(self, database: Database) -> None:
        with pytest.raises(ValueError, match="不存在"):
            SqliteBillRepository(database).confirm_bill("no-such-bill")

    def test_unknown_project_returns_empty(self, database: Database) -> None:
        repository = SqliteBillRepository(database)
        assert repository.list_bills("no-such-project") == []
        assert repository.count_bills("no-such-project") == 0
        assert repository.existing_bill_keys("no-such-project") == set()
